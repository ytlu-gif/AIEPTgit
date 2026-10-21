import os
import sys
import json
import time
import asyncio
import threading
import tempfile
import subprocess
import base64
import ctypes
import re
import wave
import uuid
import zipfile
from datetime import datetime
import aiohttp
import requests
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from docx import Document
from docx.table import Table
from docx.shared import Pt, RGBColor, Inches
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.opc.exceptions import PackageNotFoundError
from docx.text.paragraph import Paragraph
import edge_tts
from PIL import Image, ImageDraw, ImageFont, ImageTk

CONFIG_FILE = "config.json"
CHAT_HISTORY_FILE = "chat_history.json"
APP_TITLE = "AI互動式英語練習工具"
APP_COLORS = {
    "background": "#F2F6FA",
    "surface": "#FFFFFF",
    "navy": "#173B57",
    "blue": "#246A8D",
    "teal": "#238B83",
    "text": "#243746",
    "muted": "#5D7080",
    "border": "#D4E0E8",
}
AI_PROVIDERS = {
    "Gemini": "GEMINI_API_KEY",
    "OpenAI": "OPENAI_API_KEY",
    "Claude": "ANTHROPIC_API_KEY",
}
GEMINI_MODEL = "gemini-3.8-flash"


def build_learning_prompt(topic, keywords, scenario):
    keyword_list = ", ".join(keywords)
    return f"""
你是一位專業的英語教學專家。請針對「{scenario}」情境、主題「{topic}」以及目標關鍵字「{keyword_list}」，製作結構化英語學習卡內容。
逗號分隔的每一項都是一個完整關鍵字；其中可能包含空格的片語，請勿將片語拆成多個關鍵字。
使用者輸入的關鍵字可能是繁體中文、英文或中英混合。請理解每個原始關鍵字的意思，並在 words 陣列中將「word」填入最自然、常用且符合主題的英文單字或片語；不可把中文原文直接當成英文 word，也不可漏掉任何一個關鍵字。若原文是中文，請在該詞條的「pos」中文釋義清楚呈現原文含義，並以英文詞條撰寫文章、例句、搭配詞及用法。若原文已有英文，保留其正確英文詞彙或片語。
請嚴格輸出符合規範的純 JSON 格式資料（不要加任何 Markdown 文字說明或 ```json 標記）：

{{
  "title": "英文標題",
  "title_zh": "中文標題",
  "article": {{
    "en": [
      "第一段英文文章 (包含目標單字)",
      "第二段英文文章",
      "第三段英文文章"
    ],
    "zh": [
      "第一段中文翻譯",
      "第二段中文翻譯",
      "第三段中文翻譯"
    ]
  }},
  "words": [
    {{
      "word": "對應的英文單字或片語",
      "kk": "KK 音標符號（不含方括號）；片語逐詞標註並以空格分隔",
      "pos": [{{"pos": "n.", "zh": "詞性與中文釋義"}}],
      "forms": "詞形變化與相關詞說明",
      "usage": "用法與介系詞搭配說明",
      "examples": [
        {{"en": "英文例句1", "zh": "中文例句1翻譯"}},
        {{"en": "英文例句2", "zh": "中文例句2翻譯"}}
      ],
      "synonyms": ["同義字1", "同義字2"],
      "collocations": ["常用搭配詞1", "常用搭配詞2"],
      "listening": [
        {{"en": "情境英文問答", "zh": "情境中文翻譯"}}
      ]
    }}
  ]
}}

要求：
1. 文章 (article.en) 全文長度必須介於 200～300 字之間，分成 3 段。
2. 文章以「{scenario}」為背景，適合一般英語學習者，內容自然且符合該情境，並完整包含每個關鍵字所對應的英文單字或片語。
3. words 陣列須為每個輸入關鍵字各提供一筆精解；「word」一律使用對應的英文單字或片語，不可將片語拆開；每筆都必須提供正確的美式 KK 音標符號（不含方括號），而非 IPA，片語須逐詞標註。
4. 輸出必須為標準可被 json.loads 解析的純 JSON 字串。
"""


def parse_ai_json(text, provider):
    if not isinstance(text, str):
        raise RuntimeError(f"{provider} 回傳的內容不是文字格式。")
    text = text.strip()
    fenced_json = re.fullmatch(
        r"```(?:json)?\s*(.*?)\s*```", text, flags=re.IGNORECASE | re.DOTALL
    )
    if fenced_json:
        text = fenced_json.group(1)

    try:
        data = json.loads(text)
    except json.JSONDecodeError as error:
        raise RuntimeError(
            f"{provider} 回傳的內容不是有效 JSON "
            f"(第 {error.lineno} 行，第 {error.colno} 欄)。"
        ) from error

    if not isinstance(data, dict) or not isinstance(data.get("words"), list):
        raise RuntimeError(f"{provider} 回傳內容缺少有效的 words 清單。")
    for word_info in data["words"]:
        if (
            not isinstance(word_info, dict)
            or not isinstance(word_info.get("kk"), str)
            or not word_info["kk"].strip()
        ):
            word = (
                word_info.get("word", "未知單字")
                if isinstance(word_info, dict)
                else "未知單字"
            )
            raise RuntimeError(f"{provider} 未提供「{word}」的 KK 音標，請重試。")
    return data


def parse_target_keywords(raw_keywords):
    return [keyword.strip() for keyword in raw_keywords.split(",") if keyword.strip()]


def append_keyword_transcription(existing_text, transcription):
    existing_text = existing_text.strip().rstrip(",，;； ")
    transcription = transcription.strip().lstrip(",，;； ")
    if not existing_text:
        return transcription
    if not transcription:
        return existing_text
    return f"{existing_text}, {transcription}"


def build_keyword_pattern(keywords):
    if not keywords:
        return None

    alternatives = [
        re.escape(keyword).replace(r"\ ", r"\s+")
        for keyword in sorted(keywords, key=len, reverse=True)
    ]
    return re.compile(
        r"(?<!\w)(?:" + "|".join(alternatives) + r")(?!\w)",
        re.IGNORECASE
    )


def add_highlighted_text(paragraph, text, keyword_pattern):
    runs = []
    position = 0

    if keyword_pattern:
        for match in keyword_pattern.finditer(text):
            if match.start() > position:
                runs.append(paragraph.add_run(text[position:match.start()]))
            highlighted_run = paragraph.add_run(match.group())
            highlighted_run.font.color.rgb = RGBColor(0, 112, 192)
            runs.append(highlighted_run)
            position = match.end()

    if position < len(text) or not runs:
        runs.append(paragraph.add_run(text[position:]))

    return runs


def set_cell_highlighted_text(cell, text, keyword_pattern):
    return add_highlighted_text(cell.paragraphs[0], text, keyword_pattern)


def set_table_borders(table):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "8")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "7F8C8D")
        borders.append(border)
    table._tbl.tblPr.append(borders)


class AppLogger:
    def __init__(self, text_widget):
        self.text_widget = text_widget

    def log(self, message):
        def _append():
            self.text_widget.insert(tk.END, message + "\n")
            self.text_widget.see(tk.END)
        self.text_widget.after(0, _append)


class FlashcardApp:
    def __init__(self, root):
        self.root = root
        self.root.title(APP_TITLE)
        self.root.geometry("980x900")
        self.root.minsize(760, 700)
        self.root.resizable(True, True)
        self.root.protocol("WM_DELETE_WINDOW", self.close_application)
        self.root.configure(bg=APP_COLORS["background"])

        # 載入本機設定
        self.config = self.load_config()
        self.unsaved_api_keys = {}
        self.settings_unsaved_api_keys = {}
        self.current_provider = "Gemini"
        self.speech_provider_var = tk.StringVar(
            value=self.config.get("speech_provider", "Windows 本機")
        )
        if self.speech_provider_var.get() not in (
            "Windows 本機",
            "Gemini 雲端",
            "OpenAI 雲端",
        ):
            self.speech_provider_var.set("Windows 本機")
        self.voice_var = tk.StringVar(
            value=self.config.get(
                "voice", "en-US-JennyNeural (美式女聲)"
            )
        )
        if self.voice_var.get() not in (
            "en-US-JennyNeural (美式女聲)",
            "en-US-GuyNeural (美式男聲)",
            "en-GB-SoniaNeural (英式女聲)",
            "en-GB-RyanNeural (英式男聲)",
        ):
            self.voice_var.set("en-US-JennyNeural (美式女聲)")
        self.output_dir_var = tk.StringVar(
            value=self.config.get("output_dir", os.getcwd())
        )
        self.settings_window = None
        self.practice_context = None
        self.practice_window = None
        self.practice_messages = []
        self.practice_history_id = None
        self.practice_request_token = 0
        self.chat_history = []
        self.chat_history_error = None
        self.speech_input_target = None
        self.active_speech_button = None
        self.speech_button_default_text = ""
        self.generated_docx_path = None
        self.generated_audio_path = None
        self.audio_mci = None
        self.audio_device_alias = None
        self.audio_poll_id = None
        try:
            self.chat_history = self.load_chat_history()
        except (OSError, ValueError, TypeError) as error:
            self.chat_history_error = str(error)
        self.speech_recording_event = None

        # 設定 Logo 圖示
        self.setup_app_icon()

        self.setup_ui()

    def load_config(self):
        """讀取本機設定檔 config.json"""
        if os.path.exists(CONFIG_FILE):
            try:
                with open(CONFIG_FILE, "r", encoding="utf-8") as f:
                    return json.load(f)
            except Exception:
                pass
        return {}

    def save_config(self, api_key, provider=None):
        provider = provider or self.provider_combo.get()
        updated_config = dict(self.config)
        api_keys = dict(updated_config.get("api_keys", {}))
        api_keys[provider] = api_key
        updated_config["api_keys"] = api_keys
        if provider == "Gemini":
            updated_config["api_key"] = api_key
        try:
            with open(CONFIG_FILE, "w", encoding="utf-8") as f:
                json.dump(updated_config, f, ensure_ascii=False, indent=2)
            self.config = updated_config
            return True
        except OSError as e:
            print(f"儲存設定檔失敗: {e}")
            return False

    def delete_config(self, provider):
        updated_config = dict(self.config)
        api_keys = dict(updated_config.get("api_keys", {}))
        api_keys.pop(provider, None)
        updated_config["api_keys"] = api_keys
        if provider == "Gemini":
            updated_config.pop("api_key", None)
        if not api_keys and not updated_config.get("api_key"):
            updated_config.pop("api_keys", None)
        try:
            if updated_config:
                with open(CONFIG_FILE, "w", encoding="utf-8") as f:
                    json.dump(updated_config, f, ensure_ascii=False, indent=2)
            elif os.path.exists(CONFIG_FILE):
                os.remove(CONFIG_FILE)
            self.config = updated_config
            return True
        except OSError as e:
            print(f"刪除設定檔失敗: {e}")
            return False

    def manual_save_key(self):
        provider = self.settings_provider_combo.get()
        key = self.settings_key_entry.get().strip()
        if not key:
            messagebox.showwarning("提示", "API Key 欄位為空，無法儲存！")
            return
        if self.save_config(key, provider):
            self.settings_unsaved_api_keys[provider] = key
            if self.provider_combo.get() == provider:
                self.api_key_entry.delete(0, tk.END)
                self.api_key_entry.insert(0, key)
            messagebox.showinfo(
                "成功", f"{provider} API Key 已成功儲存至本機！"
            )
        else:
            messagebox.showerror("失敗", "儲存 API Key 時發生錯誤。")

    def manual_delete_key(self):
        provider = self.settings_provider_combo.get()
        if messagebox.askyesno(
            "確認刪除", f"確定要清空並刪除本機儲存的 {provider} API Key 嗎？"
        ):
            self.settings_key_entry.delete(0, tk.END)
            if self.delete_config(provider):
                self.settings_unsaved_api_keys[provider] = ""
                self.unsaved_api_keys.pop(provider, None)
                if self.provider_combo.get() == provider:
                    self.api_key_entry.delete(0, tk.END)
                    environment_key = os.getenv(AI_PROVIDERS[provider], "")
                    if environment_key:
                        self.api_key_entry.insert(0, environment_key)
                messagebox.showinfo("完成", f"已成功刪除本機儲存的 {provider} API Key！")
            else:
                messagebox.showerror("失敗", "刪除設定檔時發生錯誤。")

    def get_provider_api_key(self, provider):
        api_keys = self.config.get("api_keys", {})
        configured_key = api_keys.get(provider, "")
        if provider == "Gemini":
            configured_key = configured_key or self.config.get("api_key", "")
        return configured_key or os.getenv(AI_PROVIDERS[provider], "")

    def on_provider_change(self, event=None):
        if event is not None and event.widget is not self.provider_combo:
            return
        current_key = self.api_key_entry.get().strip()
        configured_key = self.config.get("api_keys", {}).get(self.current_provider, "")
        if self.current_provider == "Gemini":
            configured_key = configured_key or self.config.get("api_key", "")
        environment_key = os.getenv(AI_PROVIDERS[self.current_provider], "")
        if current_key and current_key not in (configured_key, environment_key):
            self.unsaved_api_keys[self.current_provider] = current_key
        else:
            self.unsaved_api_keys.pop(self.current_provider, None)
        self.current_provider = self.provider_combo.get()
        api_key = self.unsaved_api_keys.get(
            self.current_provider, self.get_provider_api_key(self.current_provider)
        )
        self.api_key_entry.delete(0, tk.END)
        if api_key:
            self.api_key_entry.insert(0, api_key)

    def setup_app_icon(self):
        """動態繪製並設定專屬應用程式 Logo 圖示"""
        try:
            size = (64, 64)
            img = Image.new("RGBA", size, (0, 0, 0, 0))
            draw = ImageDraw.Draw(img)

            draw.ellipse([2, 2, 62, 62], fill=(26, 82, 118))
            draw.ellipse([6, 6, 58, 58], outline=(255, 255, 255), width=2)

            try:
                font = ImageFont.truetype("arial.ttf", 32)
            except IOError:
                font = ImageFont.load_default()

            draw.text((21, 10), "A", fill=(255, 255, 255), font=font)
            draw.arc([38, 38, 52, 52], start=270, end=90, fill=(241, 196, 15), width=3)

            self.icon_photo = ImageTk.PhotoImage(img)
            self.root.iconphoto(False, self.icon_photo)

        except Exception as e:
            print(f"Logo 設定跳過或失敗: {e}")

    def setup_ui(self):
        style = ttk.Style(self.root)
        if "clam" in style.theme_names():
            style.theme_use("clam")
        style.configure("App.TFrame", background=APP_COLORS["background"])
        style.configure(
            "Card.TLabelframe",
            background=APP_COLORS["surface"],
            bordercolor=APP_COLORS["border"],
            relief="solid",
        )
        style.configure(
            "Card.TLabelframe.Label",
            background=APP_COLORS["surface"],
            foreground=APP_COLORS["blue"],
            font=("Microsoft JhengHei", 10, "bold"),
        )
        style.configure(
            "Accent.TButton",
            background=APP_COLORS["teal"],
            foreground="#FFFFFF",
            font=("Microsoft JhengHei", 10, "bold"),
            padding=(10, 7),
        )
        style.map(
            "Accent.TButton",
            background=[
                ("disabled", "#AAB9C2"),
                ("active", APP_COLORS["blue"]),
                ("!disabled", APP_COLORS["teal"]),
            ],
            foreground=[("disabled", "#F2F6FA"), ("!disabled", "#FFFFFF")],
        )
        style.configure(
            "Secondary.TButton",
            background="#E4EEF4",
            foreground=APP_COLORS["navy"],
            padding=(8, 5),
        )
        style.map(
            "Secondary.TButton",
            background=[("active", "#D3E4EC"), ("!disabled", "#E4EEF4")],
        )
        style.configure(
            "History.Treeview",
            background=APP_COLORS["surface"],
            fieldbackground=APP_COLORS["surface"],
            foreground=APP_COLORS["text"],
            rowheight=28,
            font=("Microsoft JhengHei", 9),
        )
        style.configure(
            "History.Treeview.Heading",
            background="#E4EEF4",
            foreground=APP_COLORS["navy"],
            font=("Microsoft JhengHei", 9, "bold"),
        )

        menu_bar = tk.Menu(self.root)
        file_menu = tk.Menu(menu_bar, tearoff=False)
        file_menu.add_command(label="檔案管理...", command=self.open_file_manager)
        file_menu.add_separator()
        file_menu.add_command(label="結束", command=self.root.destroy)
        menu_bar.add_cascade(label="檔案", menu=file_menu)

        settings_menu = tk.Menu(menu_bar, tearoff=False)
        settings_menu.add_command(label="偏好設定...", command=self.open_settings_window)
        menu_bar.add_cascade(label="設定", menu=settings_menu)

        self.tools_menu = tk.Menu(menu_bar, tearoff=False)
        self.tools_menu.add_command(
            label="AI 英文對話練習",
            command=self.open_practice_window,
            state=tk.DISABLED,
        )
        self.practice_menu_index = 0
        self.tools_menu.add_command(
            label="對話紀錄...", command=self.open_conversation_history
        )
        menu_bar.add_cascade(label="工具", menu=self.tools_menu)

        help_menu = tk.Menu(menu_bar, tearoff=False)
        help_menu.add_command(label="使用說明", command=self.open_usage_window)
        help_menu.add_command(label="版權資訊", command=self.open_copyright_window)
        menu_bar.add_cascade(label="說明", menu=help_menu)
        self.root.config(menu=menu_bar)

        # 標題區
        title_frame = tk.Frame(
            self.root, bg=APP_COLORS["navy"], padx=18, pady=13
        )
        title_frame.pack(fill=tk.X)

        tk.Label(
            title_frame,
            text=APP_TITLE,
            font=("Microsoft JhengHei", 18, "bold"),
            fg="#FFFFFF",
            bg=APP_COLORS["navy"],
        ).pack(anchor=tk.W)
        tk.Label(
            title_frame,
            text="雙語學習卡・語音輸入・AI 英語對話練習",
            font=("Microsoft JhengHei", 10),
            fg="#D7EAF3",
            bg=APP_COLORS["navy"],
        ).pack(anchor=tk.W, pady=(4, 0))

        # 設定參數區
        param_frame = ttk.LabelFrame(
            self.root,
            text="學習內容與 AI 服務",
            padding=15,
            style="Card.TLabelframe",
        )
        param_frame.pack(fill=tk.X, padx=15, pady=10)

        # API Key 欄位區（搭配儲存與刪除按鈕）
        ttk.Label(param_frame, text="AI 服務:").grid(row=0, column=0, sticky=tk.W, pady=5)
        self.provider_combo = ttk.Combobox(
            param_frame, width=47, state="readonly", values=list(AI_PROVIDERS)
        )
        self.provider_combo.set(self.current_provider)
        self.provider_combo.grid(row=0, column=1, columnspan=2, sticky=tk.W, pady=5)
        self.provider_combo.bind("<<ComboboxSelected>>", self.on_provider_change)

        ttk.Label(param_frame, text="API Key:").grid(row=1, column=0, sticky=tk.W, pady=5)

        self.api_key_entry = ttk.Entry(param_frame, width=36, show="*")
        self.api_key_entry.grid(row=1, column=1, sticky=tk.W, pady=5, padx=(0, 5))

        saved_key = self.get_provider_api_key(self.current_provider)
        if saved_key:
            self.api_key_entry.insert(0, saved_key)

        ttk.Label(param_frame, text="學習情境:").grid(row=2, column=0, sticky=tk.W, pady=5)
        self.scenario_combo = ttk.Combobox(
            param_frame, width=47, state="readonly",
            values=("商業", "學術", "日常生活", "旅遊", "職場溝通", "綜合")
        )
        self.scenario_combo.set("商業")
        self.scenario_combo.grid(row=2, column=1, columnspan=2, sticky=tk.W, pady=5)

        # 主題
        ttk.Label(param_frame, text="學習單元主題:").grid(row=3, column=0, sticky=tk.W, pady=5)
        self.topic_entry = ttk.Entry(param_frame, width=50)
        self.topic_entry.grid(row=3, column=1, columnspan=2, sticky=tk.W, pady=5)
        self.topic_entry.insert(0, "Workplace Communication")

        # 目標單字或片語
        ttk.Label(
            param_frame, text="關鍵字（中文／英文，逗號分隔）:"
        ).grid(row=4, column=0, sticky=tk.W, pady=5)
        self.words_entry = ttk.Entry(param_frame, width=38)
        self.words_entry.grid(row=4, column=1, sticky=tk.W, pady=5)
        self.words_entry.insert(0, "schedule, follow up, deadline")
        self.speech_btn = ttk.Button(
            param_frame,
            text="🎙 語音輸入",
            command=self.start_speech_input,
        )
        self.speech_btn.grid(row=4, column=2, sticky=tk.W, padx=2, pady=5)

        # 操作按鈕
        btn_frame = ttk.Frame(self.root, padding=(15, 4), style="App.TFrame")
        btn_frame.pack(fill=tk.X, padx=15)
        self.start_btn = ttk.Button(
            btn_frame,
            text="開始生成學習卡與語音檔",
            command=self.start_generation_thread,
            style="Accent.TButton",
        )
        self.start_btn.pack(fill=tk.X, ipady=8)
        self.practice_btn = ttk.Button(
            btn_frame,
            text="開啟 AI 英文對話練習",
            command=self.open_practice_window,
            state=tk.DISABLED,
            style="Secondary.TButton",
        )
        self.practice_btn.pack(fill=tk.X, ipady=5, pady=(6, 0))

        preview_frame = ttk.LabelFrame(
            self.root,
            text="學習卡預覽與語音播放",
            padding=10,
            style="Card.TLabelframe",
        )
        preview_frame.pack(fill=tk.BOTH, expand=True, padx=15, pady=(8, 4))
        preview_frame.columnconfigure(0, weight=1)
        preview_frame.rowconfigure(1, weight=1)

        preview_toolbar = ttk.Frame(preview_frame, style="App.TFrame")
        preview_toolbar.grid(row=0, column=0, columnspan=2, sticky=tk.EW, pady=(0, 6))
        self.preview_title_var = tk.StringVar(value="尚未生成學習卡")
        ttk.Label(
            preview_toolbar,
            textvariable=self.preview_title_var,
            font=("Microsoft JhengHei", 10, "bold"),
        ).pack(side=tk.LEFT, fill=tk.X, expand=True)
        self.open_card_btn = ttk.Button(
            preview_toolbar,
            text="開啟 Word 學習卡",
            command=self.open_generated_card,
            state=tk.DISABLED,
            style="Secondary.TButton",
        )
        self.open_card_btn.pack(side=tk.RIGHT)
        self.import_audio_btn = ttk.Button(
            preview_toolbar,
            text="匯入歷史語音",
            command=self.import_historical_audio,
            style="Secondary.TButton",
        )
        self.import_audio_btn.pack(side=tk.RIGHT, padx=(0, 6))
        self.import_card_btn = ttk.Button(
            preview_toolbar,
            text="匯入歷史學習卡",
            command=self.import_historical_card,
            style="Secondary.TButton",
        )
        self.import_card_btn.pack(side=tk.RIGHT, padx=(0, 6))

        self.card_preview_text = tk.Text(
            preview_frame,
            wrap=tk.WORD,
            state=tk.DISABLED,
            font=("Microsoft JhengHei", 10),
            bg="#FBFDFF",
            fg=APP_COLORS["text"],
            relief=tk.FLAT,
            padx=10,
            pady=8,
        )
        self.card_preview_text.grid(row=1, column=0, sticky=tk.NSEW)
        preview_scrollbar = ttk.Scrollbar(
            preview_frame, command=self.card_preview_text.yview
        )
        preview_scrollbar.grid(row=1, column=1, sticky=tk.NS)
        self.card_preview_text.configure(yscrollcommand=preview_scrollbar.set)
        self.card_preview_text.tag_configure(
            "preview-title",
            font=("Microsoft JhengHei", 15, "bold"),
            foreground=APP_COLORS["navy"],
            spacing3=4,
        )
        self.card_preview_text.tag_configure(
            "preview-subtitle",
            font=("Microsoft JhengHei", 11, "bold"),
            foreground=APP_COLORS["blue"],
            spacing1=10,
            spacing3=3,
        )
        self.card_preview_text.tag_configure(
            "preview-word",
            font=("Microsoft JhengHei", 11, "bold"),
            foreground=APP_COLORS["teal"],
            spacing1=7,
        )
        self.card_preview_text.tag_configure(
            "muted", foreground=APP_COLORS["muted"]
        )
        self.card_preview_text.insert(
            tk.END,
            "生成完成後，學習卡內容會顯示在這裡；也可匯入歷史 Word 學習卡及 MP3/WAV 語音檔。",
        )
        self.card_preview_text.configure(state=tk.DISABLED)

        audio_controls = ttk.Frame(preview_frame, style="App.TFrame")
        audio_controls.grid(
            row=2, column=0, columnspan=2, sticky=tk.EW, pady=(7, 0)
        )
        self.play_audio_btn = ttk.Button(
            audio_controls,
            text="▶ 播放",
            command=self.play_generated_audio,
            state=tk.DISABLED,
            style="Accent.TButton",
        )
        self.play_audio_btn.pack(side=tk.LEFT)
        self.pause_audio_btn = ttk.Button(
            audio_controls,
            text="⏸ 暫停",
            command=self.pause_generated_audio,
            state=tk.DISABLED,
            style="Secondary.TButton",
        )
        self.pause_audio_btn.pack(side=tk.LEFT, padx=(6, 0))
        self.stop_audio_btn = ttk.Button(
            audio_controls,
            text="■ 停止",
            command=self.stop_generated_audio,
            state=tk.DISABLED,
            style="Secondary.TButton",
        )
        self.stop_audio_btn.pack(side=tk.LEFT, padx=(6, 10))
        self.audio_status_var = tk.StringVar(value="尚無語音檔")
        ttk.Label(
            audio_controls,
            textvariable=self.audio_status_var,
            foreground=APP_COLORS["muted"],
        ).pack(side=tk.LEFT, fill=tk.X, expand=True)

        # 處理進度記錄
        log_frame = ttk.LabelFrame(
            self.root,
            text="處理進度",
            padding=8,
            style="Card.TLabelframe",
        )
        log_frame.pack(fill=tk.X, padx=15, pady=(4, 10))
        log_frame.configure(height=125)
        log_frame.pack_propagate(False)

        self.log_text = tk.Text(log_frame, wrap=tk.WORD, font=("Consolas", 9))
        self.log_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar = ttk.Scrollbar(log_frame, command=self.log_text.yview)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.log_text.config(yscrollcommand=scrollbar.set)

        self.logger = AppLogger(self.log_text)

    def open_usage_window(self):
        window = tk.Toplevel(self.root)
        window.title("使用說明")
        window.geometry("620x560")
        window.minsize(520, 420)
        window.transient(self.root)

        guide = (
            "AI互動式英語練習工具｜使用說明\n\n"
            "一、生成學習卡\n"
            "1. 選擇 AI 服務並輸入 API Key。可至「設定 → 偏好設定」管理已儲存的"
            "金鑰，或以環境變數提供。\n"
            "2. 選擇學習情境，輸入主題與以逗號分隔的目標單字或片語；"
            "關鍵字可用繁體中文或英文，AI 會轉成對應英文詞條。\n"
            "3. 可選擇語音辨識方式；雲端辨識會在錄音前詢問是否同意傳送音訊，"
            "並可能產生 API 費用。\n"
            "4. 在「設定 → 偏好設定」選擇朗讀聲音與輸出位置，再按"
            "「開始生成學習卡與語音檔」。\n"
            "5. 程式會輸出 Word 學習卡與文章語音檔；若 MP3 服務無法使用，"
            "Windows 可改輸出 WAV。生成完成後，主視窗會顯示學習卡預覽，"
            "並提供語音播放、暫停、繼續及停止控制；「開啟 Word 學習卡」可查看完整文件。\n\n"
            "若要查看舊檔，按「匯入歷史學習卡」選取 .docx；"
            "若同資料夾中有相同主題名稱的 _article.mp3 或 _article.wav，"
            "程式會自動配對，也可按「匯入歷史語音」自行選取 MP3/WAV。\n\n"
            "二、語音輸入\n"
            "「Windows 本機」在目標字詞欄使用繁體中文辨識器，在對話練習欄使用英文辨識器；"
            "請先在 Windows 安裝對應語言的語音辨識功能。"
            "「Gemini 雲端」與「OpenAI 雲端」會自動辨識繁體中文或英文，"
            "需有效的對應 API Key 與網路連線。\n"
            "辨識結果會填入目前使用的輸入欄；目標字詞請以逗號分隔。\n\n"
            "三、AI 對話練習\n"
            "成功生成學習卡後，按「開啟 AI 英文對話練習」。AI 會依據本次學習卡的"
            "主題、情境與目標字詞進行文字對話；可點選提示句帶入輸入欄，"
            "或按輸入欄旁的麥克風說英文。\n\n"
            "「工具 → 對話紀錄」可選取舊對話接續練習，對話內容會自動儲存在本機。\n\n"
            "四、檔案管理\n"
            "按「檔案 → 檔案管理」可瀏覽資料夾、重新整理清單、"
            "開啟選取檔案或在檔案總管中開啟資料夾。"
        )
        text = tk.Text(
            window, wrap=tk.WORD, font=("Microsoft JhengHei", 10), padx=12, pady=12
        )
        scrollbar = ttk.Scrollbar(window, command=text.yview)
        text.configure(yscrollcommand=scrollbar.set)
        text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(10, 0), pady=10)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y, padx=(0, 10), pady=10)
        text.insert("1.0", guide)
        text.configure(state=tk.DISABLED)

    def open_copyright_window(self):
        window = tk.Toplevel(self.root)
        window.title("版權資訊")
        window.geometry("460x300")
        window.resizable(False, False)
        window.transient(self.root)

        content = ttk.Frame(window, padding=24)
        content.pack(fill=tk.BOTH, expand=True)
        ttk.Label(
            content, text="AI互動式英語練習工具", font=("Microsoft JhengHei", 15, "bold")
        ).pack(anchor=tk.W, pady=(0, 14))
        ttk.Label(content, text="版本：2.7.0", font=("Microsoft JhengHei", 10)).pack(
            anchor=tk.W, pady=4
        )
        ttk.Label(
            content,
            text="發行者：Yu-Tsang Lu at CMU",
            font=("Microsoft JhengHei", 10),
        ).pack(anchor=tk.W, pady=4)
        ttk.Label(
            content,
            text=f"著作權 © {datetime.now().year} Yu-Tsang Lu at CMU",
            font=("Microsoft JhengHei", 10),
        ).pack(anchor=tk.W, pady=4)
        ttk.Label(
            content,
            text=(
                "本程式使用第三方服務與開源套件。AI 服務、語音辨識及語音合成"
                "各自受其服務條款與隱私政策約束。"
            ),
            wraplength=400,
            justify=tk.LEFT,
            font=("Microsoft JhengHei", 9),
        ).pack(anchor=tk.W, pady=(12, 4))
        ttk.Button(content, text="關閉", command=window.destroy).pack(
            anchor=tk.E, pady=(12, 0)
        )

    def open_file_manager(self):
        if (
            hasattr(self, "file_manager_window")
            and self.file_manager_window.winfo_exists()
        ):
            self.file_manager_window.lift()
            return

        window = tk.Toplevel(self.root)
        window.title("檔案管理")
        window.geometry("760x500")
        window.minsize(600, 360)
        window.transient(self.root)
        window.columnconfigure(0, weight=1)
        window.rowconfigure(1, weight=1)
        self.file_manager_window = window
        self.file_manager_paths = {}

        folder_frame = ttk.Frame(window, padding=(12, 12, 12, 8))
        folder_frame.grid(row=0, column=0, sticky=tk.EW)
        folder_frame.columnconfigure(1, weight=1)
        ttk.Label(folder_frame, text="資料夾：").grid(row=0, column=0, sticky=tk.W)
        self.file_manager_folder = tk.StringVar(
            value=self.output_dir_var.get().strip() or os.getcwd()
        )
        ttk.Entry(
            folder_frame, textvariable=self.file_manager_folder, state="readonly"
        ).grid(row=0, column=1, sticky=tk.EW, padx=6)
        ttk.Button(
            folder_frame, text="選擇資料夾...", command=self.choose_file_manager_folder
        ).grid(row=0, column=2, padx=(0, 6))
        ttk.Button(
            folder_frame, text="開啟資料夾", command=self.open_file_manager_folder
        ).grid(row=0, column=3)

        list_frame = ttk.Frame(window, padding=(12, 0, 12, 8))
        list_frame.grid(row=1, column=0, sticky=tk.NSEW)
        list_frame.columnconfigure(0, weight=1)
        list_frame.rowconfigure(0, weight=1)
        columns = ("type", "size", "modified")
        self.file_manager_tree = ttk.Treeview(
            list_frame, columns=columns, show="tree headings", selectmode="browse"
        )
        self.file_manager_tree.heading("#0", text="名稱")
        self.file_manager_tree.heading("type", text="類型")
        self.file_manager_tree.heading("size", text="大小")
        self.file_manager_tree.heading("modified", text="修改時間")
        self.file_manager_tree.column("#0", width=300, anchor=tk.W)
        self.file_manager_tree.column("type", width=100, anchor=tk.W)
        self.file_manager_tree.column("size", width=100, anchor=tk.E)
        self.file_manager_tree.column("modified", width=150, anchor=tk.W)
        tree_scrollbar = ttk.Scrollbar(
            list_frame, command=self.file_manager_tree.yview
        )
        self.file_manager_tree.configure(yscrollcommand=tree_scrollbar.set)
        self.file_manager_tree.grid(row=0, column=0, sticky=tk.NSEW)
        tree_scrollbar.grid(row=0, column=1, sticky=tk.NS)
        self.file_manager_tree.bind("<Double-1>", self.open_selected_file)

        actions = ttk.Frame(window, padding=(12, 0, 12, 12))
        actions.grid(row=2, column=0, sticky=tk.EW)
        ttk.Button(
            actions, text="重新整理", command=self.refresh_file_manager
        ).pack(side=tk.LEFT)
        ttk.Button(
            actions, text="開啟選取項目", command=self.open_selected_file
        ).pack(side=tk.RIGHT)
        self.refresh_file_manager()

    def choose_file_manager_folder(self):
        folder = filedialog.askdirectory(
            initialdir=self.file_manager_folder.get() or os.getcwd(),
            parent=self.file_manager_window,
        )
        if folder:
            self.file_manager_folder.set(folder)
            self.refresh_file_manager()

    def refresh_file_manager(self):
        folder = self.file_manager_folder.get()
        existing_items = self.file_manager_tree.get_children()
        if existing_items:
            self.file_manager_tree.delete(*existing_items)
        self.file_manager_paths.clear()
        if not os.path.isdir(folder):
            messagebox.showerror(
                "資料夾不存在", f"找不到輸出資料夾：\n{folder}",
                parent=self.file_manager_window,
            )
            return
        try:
            entries = sorted(
                os.scandir(folder),
                key=lambda entry: (not entry.is_dir(), entry.name.casefold()),
            )
            for entry in entries:
                try:
                    metadata = entry.stat()
                    is_directory = entry.is_dir()
                    item_id = self.file_manager_tree.insert(
                        "",
                        tk.END,
                        text=entry.name,
                        values=(
                            "資料夾" if is_directory else (os.path.splitext(entry.name)[1] or "檔案"),
                            "" if is_directory else self.format_file_size(metadata.st_size),
                            datetime.fromtimestamp(metadata.st_mtime).strftime(
                                "%Y-%m-%d %H:%M"
                            ),
                        ),
                    )
                    self.file_manager_paths[item_id] = entry.path
                except OSError as error:
                    self.logger.log(f"無法讀取檔案資訊 {entry.path}: {error}")
        except OSError as error:
            messagebox.showerror(
                "讀取失敗",
                f"無法讀取資料夾內容：\n{error}",
                parent=self.file_manager_window,
            )

    @staticmethod
    def format_file_size(size):
        value = float(size)
        for unit in ("B", "KB", "MB", "GB"):
            if value < 1024 or unit == "GB":
                return f"{value:.0f} {unit}" if unit == "B" else f"{value:.1f} {unit}"
            value /= 1024
        return f"{value:.1f} GB"

    def open_file_manager_folder(self):
        self.open_path(self.file_manager_folder.get())

    def open_selected_file(self, event=None):
        if event is not None and event.widget is not self.file_manager_tree:
            return
        selected = self.file_manager_tree.selection()
        if not selected:
            messagebox.showinfo(
                "提示", "請先選擇要開啟的檔案或資料夾。",
                parent=self.file_manager_window,
            )
            return
        path = self.file_manager_paths.get(selected[0])
        if path:
            self.open_path(path)

    def open_path(self, path):
        if not os.path.exists(path):
            messagebox.showerror("找不到項目", f"檔案或資料夾不存在：\n{path}")
            return
        try:
            if sys.platform == "win32":
                os.startfile(path)
            elif sys.platform == "darwin":
                subprocess.Popen(["open", path])
            else:
                subprocess.Popen(["xdg-open", path])
        except (OSError, subprocess.SubprocessError) as error:
            messagebox.showerror("開啟失敗", f"無法開啟項目：\n{error}")

    def open_settings_window(self):
        if self.settings_window and self.settings_window.winfo_exists():
            self.settings_window.lift()
            return

        window = tk.Toplevel(self.root)
        self.settings_window = window
        window.title("偏好設定")
        window.geometry("560x470")
        window.resizable(False, False)
        window.transient(self.root)
        window.protocol("WM_DELETE_WINDOW", window.destroy)

        frame = ttk.Frame(window, padding=18)
        frame.pack(fill=tk.BOTH, expand=True)
        frame.columnconfigure(1, weight=1)
        self.settings_speech_var = tk.StringVar(value=self.speech_provider_var.get())
        self.settings_voice_var = tk.StringVar(value=self.voice_var.get())
        self.settings_output_dir_var = tk.StringVar(value=self.output_dir_var.get())

        ttk.Label(frame, text="語音辨識方式：").grid(
            row=0, column=0, sticky=tk.W, pady=7
        )
        self.settings_speech_combo = ttk.Combobox(
            frame,
            textvariable=self.settings_speech_var,
            values=("Windows 本機", "Gemini 雲端", "OpenAI 雲端"),
            state="readonly",
            width=38,
        )
        self.settings_speech_combo.grid(
            row=0, column=1, sticky=tk.EW, pady=7
        )

        ttk.Label(frame, text="語音朗讀發音人：").grid(
            row=1, column=0, sticky=tk.W, pady=7
        )
        self.settings_voice_combo = ttk.Combobox(
            frame,
            textvariable=self.settings_voice_var,
            values=(
                "en-US-JennyNeural (美式女聲)",
                "en-US-GuyNeural (美式男聲)",
                "en-GB-SoniaNeural (英式女聲)",
                "en-GB-RyanNeural (英式男聲)",
            ),
            state="readonly",
            width=38,
        )
        if self.settings_voice_var.get() not in self.settings_voice_combo["values"]:
            self.settings_voice_var.set("en-US-JennyNeural (美式女聲)")
        self.settings_voice_combo.grid(row=1, column=1, sticky=tk.EW, pady=7)

        ttk.Label(frame, text="輸出資料夾：").grid(
            row=2, column=0, sticky=tk.W, pady=7
        )
        output_frame = ttk.Frame(frame)
        output_frame.grid(row=2, column=1, sticky=tk.EW, pady=7)
        output_frame.columnconfigure(0, weight=1)
        ttk.Entry(output_frame, textvariable=self.settings_output_dir_var).grid(
            row=0, column=0, sticky=tk.EW, padx=(0, 6)
        )
        ttk.Button(
            output_frame, text="瀏覽...", command=self.choose_settings_output_folder
        ).grid(row=0, column=1)

        ttk.Separator(frame).grid(
            row=3, column=0, columnspan=2, sticky=tk.EW, pady=12
        )
        ttk.Label(
            frame, text="API Key 管理", font=("Microsoft JhengHei", 11, "bold")
        ).grid(row=4, column=0, columnspan=2, sticky=tk.W, pady=(0, 5))
        ttk.Label(frame, text="AI 服務：").grid(
            row=5, column=0, sticky=tk.W, pady=7
        )
        self.settings_provider_combo = ttk.Combobox(
            frame,
            values=list(AI_PROVIDERS),
            state="readonly",
            width=38,
        )
        self.settings_provider_combo.set(self.current_provider)
        self.settings_provider_combo.grid(
            row=5, column=1, sticky=tk.EW, pady=7
        )
        self.settings_provider_combo.bind(
            "<<ComboboxSelected>>", self.on_settings_provider_change
        )

        ttk.Label(frame, text="API Key：").grid(
            row=6, column=0, sticky=tk.W, pady=7
        )
        self.settings_key_entry = ttk.Entry(frame, show="*", width=40)
        self.settings_key_entry.grid(row=6, column=1, sticky=tk.EW, pady=7)
        current_key = self.api_key_entry.get().strip()
        configured_key = self.config.get("api_keys", {}).get(self.current_provider, "")
        if self.current_provider == "Gemini":
            configured_key = configured_key or self.config.get("api_key", "")
        environment_key = os.getenv(AI_PROVIDERS[self.current_provider], "")
        if current_key and current_key not in (configured_key, environment_key):
            self.settings_unsaved_api_keys[self.current_provider] = current_key
        self.load_settings_api_key(self.current_provider)

        api_buttons = ttk.Frame(frame)
        api_buttons.grid(row=7, column=1, sticky=tk.W, pady=(0, 8))
        ttk.Button(
            api_buttons, text="儲存 API Key", command=self.manual_save_key
        ).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(
            api_buttons, text="刪除本機 API Key", command=self.manual_delete_key
        ).pack(side=tk.LEFT)

        ttk.Label(
            frame,
            text="API Key 只會儲存在本機 config.json；環境變數設定的金鑰不會在此顯示或刪除。",
            wraplength=500,
            justify=tk.LEFT,
            font=("Microsoft JhengHei", 9),
        ).grid(row=8, column=0, columnspan=2, sticky=tk.W, pady=(4, 10))
        buttons = ttk.Frame(frame)
        buttons.grid(row=9, column=0, columnspan=2, sticky=tk.E)
        ttk.Button(
            buttons, text="儲存偏好設定", command=self.save_preferences
        ).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(buttons, text="關閉", command=window.destroy).pack(side=tk.LEFT)

    def on_settings_provider_change(self, event=None):
        if event is not None and event.widget is not self.settings_provider_combo:
            return
        previous_provider = getattr(self, "settings_current_provider", None)
        if previous_provider:
            self.settings_unsaved_api_keys[previous_provider] = (
                self.settings_key_entry.get()
            )
        provider = self.settings_provider_combo.get()
        self.settings_current_provider = provider
        self.load_settings_api_key(provider)

    def load_settings_api_key(self, provider):
        key = self.settings_unsaved_api_keys.get(provider)
        if key is None:
            key = self.config.get("api_keys", {}).get(provider, "")
            if provider == "Gemini":
                key = key or self.config.get("api_key", "")
        self.settings_key_entry.delete(0, tk.END)
        if key:
            self.settings_key_entry.insert(0, key)
        self.settings_current_provider = provider

    def choose_settings_output_folder(self):
        folder = filedialog.askdirectory(
            initialdir=self.settings_output_dir_var.get() or os.getcwd(),
            parent=self.settings_window,
        )
        if folder:
            self.settings_output_dir_var.set(folder)

    def save_preferences(self):
        output_dir = self.settings_output_dir_var.get().strip()
        if not output_dir or not os.path.isdir(output_dir):
            messagebox.showerror(
                "設定錯誤",
                "請選擇有效的輸出資料夾。",
                parent=self.settings_window,
            )
            return
        updated_config = dict(self.config)
        updated_config["speech_provider"] = self.settings_speech_var.get()
        updated_config["voice"] = self.settings_voice_var.get()
        updated_config["output_dir"] = output_dir
        try:
            with open(CONFIG_FILE, "w", encoding="utf-8") as config_file:
                json.dump(updated_config, config_file, ensure_ascii=False, indent=2)
        except OSError as error:
            messagebox.showerror(
                "設定儲存失敗",
                f"無法儲存偏好設定：\n{error}",
                parent=self.settings_window,
            )
            return
        self.config = updated_config
        self.speech_provider_var.set(self.settings_speech_var.get())
        self.voice_var.set(self.settings_voice_var.get())
        self.output_dir_var.set(output_dir)
        messagebox.showinfo(
            "設定完成", "偏好設定已儲存。", parent=self.settings_window
        )

    def start_speech_input(
        self, target_entry=None, button=None, language="zh-TW"
    ):
        target_entry = target_entry or self.words_entry
        button = button or self.speech_btn
        self.speech_input_target = target_entry
        self.active_speech_button = button
        self.speech_button_default_text = button.cget("text")
        speech_provider = self.speech_provider_var.get()
        if speech_provider == "Windows 本機":
            if sys.platform != "win32":
                messagebox.showerror(
                    "語音輸入不可用", "Windows 內建語音辨識僅支援 Windows。"
                )
                self.speech_input_target = None
                self.active_speech_button = None
                return
            button.config(state=tk.DISABLED, text="辨識中...")
            threading.Thread(
                target=self.run_speech_recognition,
                args=(language,),
                daemon=True,
            ).start()
            return

        provider = speech_provider.removesuffix(" 雲端")
        if provider not in ("Gemini", "OpenAI"):
            messagebox.showerror("語音輸入不可用", "請選擇有效的語音辨識方式。")
            self.speech_input_target = None
            self.active_speech_button = None
            return
        api_key = ""
        if self.provider_combo.get() == provider:
            api_key = self.api_key_entry.get().strip()
        api_key = api_key or self.get_provider_api_key(provider)
        if not api_key:
            messagebox.showwarning(
                "缺少 API Key",
                f"請先在「AI 服務」選擇 {provider} 並設定 API Key，"
                "再使用該服務進行雲端語音辨識。",
            )
            self.speech_input_target = None
            self.active_speech_button = None
            return
        if not messagebox.askyesno(
            "雲端語音辨識",
            f"錄音將傳送至 {provider} 進行辨識，可能產生 API 費用。"
            "可辨識繁體中文與英文。是否繼續？",
        ):
            self.speech_input_target = None
            self.active_speech_button = None
            return

        self.speech_recording_event = threading.Event()
        button.config(
            state=tk.NORMAL,
            text="■ 停止",
            command=self.stop_cloud_speech_recording,
        )
        threading.Thread(
            target=self.run_cloud_speech_recognition,
            args=(self.speech_recording_event, provider, api_key, language),
            daemon=True,
        ).start()

    def stop_cloud_speech_recording(self):
        if self.speech_recording_event:
            self.speech_recording_event.set()
            if self.active_speech_button:
                self.active_speech_button.config(state=tk.DISABLED, text="辨識中...")

    def run_speech_recognition(self, language):
        script = r"""
$ErrorActionPreference = 'Stop'
$Language = '__LANGUAGE__'
[Console]::OutputEncoding = New-Object System.Text.UTF8Encoding
try {
    Add-Type -AssemblyName System.Speech -ErrorAction Stop
    $recognizerInfo = [System.Speech.Recognition.SpeechRecognitionEngine]::InstalledRecognizers() |
        Where-Object {
            if ($Language -eq 'zh-TW') {
                $_.Culture.Name -eq 'zh-TW'
            } else {
                $_.Culture.Name -like 'en-*'
            }
        } |
        Select-Object -First 1
    if ($null -eq $recognizerInfo) {
        throw "No Windows speech recognizer is installed for $Language. Install the matching speech recognition language in Windows settings."
    }
    $recognizer = New-Object System.Speech.Recognition.SpeechRecognitionEngine($recognizerInfo)
    try {
        $recognizer.SetInputToDefaultAudioDevice()
        $grammar = New-Object System.Speech.Recognition.DictationGrammar
        $recognizer.LoadGrammar($grammar)
        $recognizer.InitialSilenceTimeout = [TimeSpan]::FromSeconds(10)
        $recognizer.BabbleTimeout = [TimeSpan]::FromSeconds(3)
        $recognizer.EndSilenceTimeout = [TimeSpan]::FromSeconds(1)
        $result = $recognizer.Recognize([TimeSpan]::FromSeconds(30))
        if ($null -eq $result -or [string]::IsNullOrWhiteSpace($result.Text)) {
            throw "No speech was recognized. Check the microphone and Windows $Language speech recognition setup."
        }
        [Console]::Out.WriteLine($result.Text)
    }
    finally {
        $recognizer.Dispose()
    }
}
catch {
    [Console]::Error.WriteLine($_.Exception.Message)
    exit 1
}
"""
        script = script.replace("__LANGUAGE__", language)
        try:
            result = subprocess.run(
                [
                    "powershell.exe",
                    "-NoProfile",
                    "-NonInteractive",
                    "-Command",
                    script,
                ],
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=40,
                check=False,
                creationflags=subprocess.CREATE_NO_WINDOW,
            )
            if result.returncode != 0:
                detail = result.stderr.strip() or "Windows 語音辨識未能完成。"
                raise RuntimeError(detail)
            transcription = result.stdout.strip()
            if not transcription:
                raise RuntimeError("沒有辨識到語音，請確認麥克風並再試一次。")
            self.root.after(
                0, lambda text=transcription: self.finish_speech_input(text, None)
            )
        except subprocess.TimeoutExpired:
            self.root.after(
                0,
                lambda: self.finish_speech_input(
                    None, "語音辨識逾時，請確認麥克風後再試一次。"
                ),
            )
        except (OSError, subprocess.SubprocessError, RuntimeError) as error:
            self.root.after(
                0,
                lambda message=str(error): self.finish_speech_input(None, message),
            )

    def run_cloud_speech_recognition(
        self, stop_event, provider, api_key, language
    ):
        try:
            import sounddevice
        except ImportError:
            self.root.after(
                0,
                lambda: self.finish_speech_input(
                    None, "雲端語音輸入需要 sounddevice 套件；請重新安裝或更新程式。"
                ),
            )
            return

        audio_path = None
        try:
            sample_rate = 16000
            block_size = 4000
            audio_chunks = []
            with sounddevice.RawInputStream(
                samplerate=sample_rate,
                channels=1,
                dtype="int16",
                blocksize=block_size,
            ) as stream:
                for _ in range(120):
                    if stop_event.is_set():
                        break
                    audio_data, _ = stream.read(block_size)
                    audio_chunks.append(bytes(audio_data))

            if not audio_chunks:
                raise RuntimeError("沒有錄到音訊，請確認麥克風後再試一次。")

            with tempfile.NamedTemporaryFile(suffix=".wav", delete=False) as audio_file:
                audio_path = audio_file.name
            with wave.open(audio_path, "wb") as wav_file:
                wav_file.setnchannels(1)
                wav_file.setsampwidth(2)
                wav_file.setframerate(sample_rate)
                wav_file.writeframes(b"".join(audio_chunks))

            transcription = self.transcribe_cloud_audio(
                provider, api_key, audio_path, language
            )
            if not transcription:
                raise RuntimeError("雲端服務沒有辨識出文字，請再試一次。")
            self.root.after(
                0, lambda text=transcription: self.finish_speech_input(text, None)
            )
        except (sounddevice.PortAudioError, OSError, RuntimeError, wave.Error) as error:
            self.root.after(
                0,
                lambda message=str(error): self.finish_speech_input(None, message),
            )
        finally:
            if audio_path and os.path.exists(audio_path):
                try:
                    os.remove(audio_path)
                except OSError as error:
                    self.logger.log(f"暫存語音檔清理失敗：{error}")

    def transcribe_cloud_audio(self, provider, api_key, audio_path, language):
        if provider == "OpenAI":
            try:
                with open(audio_path, "rb") as audio_file:
                    response = requests.post(
                        "https://api.openai.com/v1/audio/transcriptions",
                        headers={"Authorization": f"Bearer {api_key}"},
                        data={
                            "model": "whisper-1",
                            "language": "zh" if language == "zh-TW" else "en",
                            "prompt": (
                                "請以繁體中文辨識語音，保留英文專有名詞。"
                                "多個獨立單字或片語請用逗號分隔，不要翻譯或補充。"
                                if language == "zh-TW"
                                else "Transcribe the English speech verbatim. "
                                "Separate independent words or phrases with commas. "
                                "Do not translate or add explanations."
                            ),
                        },
                        files={"file": ("speech.wav", audio_file, "audio/wav")},
                        timeout=(15, 120),
                    )
            except requests.RequestException as error:
                raise RuntimeError(f"OpenAI 語音辨識連線失敗：{error}") from error
            response_data = self.parse_chat_response(response, provider)
            try:
                return response_data["text"].strip()
            except (KeyError, TypeError, AttributeError) as error:
                raise RuntimeError(
                    f"{provider} 語音辨識回傳格式不正確：{error}"
                ) from error

        if provider == "Gemini":
            try:
                with open(audio_path, "rb") as audio_file:
                    encoded_audio = base64.b64encode(audio_file.read()).decode("ascii")
                response = requests.post(
                    "https://generativelanguage.googleapis.com/v1beta/models/"
                    f"{GEMINI_MODEL}:generateContent",
                    params={"key": api_key},
                    headers={"Content-Type": "application/json"},
                    json={
                        "contents": [
                            {
                                "parts": [
                                    {
                                        "text": (
                                            (
                                                "請以繁體中文辨識錄音，保留其中的英文詞彙。"
                                                "多個獨立關鍵字請用逗號分隔。"
                                                "只輸出辨識結果，不要翻譯或補充說明。"
                                            )
                                            if language == "zh-TW"
                                            else (
                                                "Transcribe the English speech verbatim. "
                                                "Separate independent vocabulary items with "
                                                "commas. Return only the transcription without "
                                                "translation or explanations."
                                            )
                                        )
                                    },
                                    {
                                        "inlineData": {
                                            "mimeType": "audio/wav",
                                            "data": encoded_audio,
                                        }
                                    },
                                ]
                            }
                        ],
                        "generationConfig": {"temperature": 0},
                    },
                    timeout=(15, 120),
                )
            except requests.RequestException as error:
                raise RuntimeError(f"Gemini 語音辨識連線失敗：{error}") from error
            response_data = self.parse_chat_response(response, provider)
            try:
                return "".join(
                    part["text"]
                    for part in response_data["candidates"][0]["content"]["parts"]
                    if "text" in part
                ).strip()
            except (KeyError, IndexError, TypeError) as error:
                raise RuntimeError(
                    f"{provider} 語音辨識回傳格式不正確：{error}"
                ) from error

        raise ValueError(f"不支援的雲端語音服務：{provider}")

    def finish_speech_input(self, transcription, error):
        self.speech_recording_event = None
        target_entry = self.speech_input_target
        button = self.active_speech_button
        default_text = self.speech_button_default_text
        self.speech_input_target = None
        self.active_speech_button = None
        if button and button.winfo_exists():
            button.config(state=tk.NORMAL, text=default_text)
            if button is self.speech_btn:
                button.config(command=self.start_speech_input)
            else:
                button.config(
                    command=lambda: self.start_speech_input(
                        target_entry=target_entry,
                        button=button,
                        language="en-US",
                    )
                )
        if error:
            messagebox.showerror(
                "語音辨識失敗",
                f"{error}\n\n請確認已安裝對應語言的 Windows 語音辨識功能，"
                "並允許應用程式使用麥克風。",
                parent=(
                    self.practice_window
                    if button is not self.speech_btn
                    and self.practice_window
                    and self.practice_window.winfo_exists()
                    else self.root
                ),
            )
            return
        if not target_entry or not target_entry.winfo_exists():
            return
        if target_entry is self.words_entry:
            updated_text = append_keyword_transcription(
                target_entry.get(), transcription
            )
            target_entry.delete(0, tk.END)
            target_entry.insert(0, updated_text)
        else:
            current_text = target_entry.get().strip()
            target_entry.delete(0, tk.END)
            target_entry.insert(0, f"{current_text} {transcription}".strip())
            target_entry.focus_set()

    def start_generation_thread(self):
        provider = self.provider_combo.get()
        api_key = self.api_key_entry.get().strip()
        topic = self.topic_entry.get().strip()
        words = self.words_entry.get().strip()
        keywords = parse_target_keywords(words)
        save_dir = self.output_dir_var.get().strip()
        scenario = self.scenario_combo.get()

        if not api_key:
            messagebox.showwarning("提示", f"請輸入有效的 {provider} API Key！")
            return
        if not topic or not keywords:
            messagebox.showwarning("提示", "請輸入主題與至少一個目標單字或片語！")
            return
        if not scenario:
            messagebox.showwarning("提示", "請選擇學習情境！")
            return

        words = ", ".join(keywords)

        # 自動記憶並儲存 API Key 至本機 config.json
        self.save_config(api_key)
        voice_code = self.voice_var.get().split()[0]

        self.start_btn.config(state=tk.DISABLED)
        self.log_text.delete("1.0", tk.END)

        threading.Thread(
            target=self.run_process,
            args=(provider, api_key, topic, words, save_dir, scenario, voice_code),
            daemon=True
        ).start()

    def run_process(self, provider, api_key, topic, words, save_dir, scenario, voice_code):
        try:
            safe_topic = "".join([c if c.isalnum() or c in (' ', '_', '-') else '_' for c in topic]).strip().replace(" ", "_")

            self.logger.log(f"► [1/4] 正在使用 {provider} 生成結構化學習卡內容...")
            content_json = self.call_ai_api(provider, api_key, topic, words, scenario)
            card_keywords = [
                item["word"].strip()
                for item in content_json.get("words", [])
                if isinstance(item, dict)
                and isinstance(item.get("word"), str)
                and item["word"].strip()
            ] or parse_target_keywords(words)

            self.logger.log("► [2/4] 正在繪製與排版 Word 學習卡 (.docx)...")
            docx_path = self.generate_docx(
                content_json, save_dir, safe_topic, card_keywords
            )

            self.logger.log("► [3/4] 正在生成文章語音朗讀檔...")
            mp3_path = self.generate_audio(content_json, save_dir, safe_topic, voice_code)

            self.logger.log("\n✅ [4/4] 全部任務完成！")
            self.logger.log(f"📄 Word 檔儲存於: {docx_path}")
            self.logger.log(f"🔊 語音檔儲存於: {mp3_path}")
            self.root.after(
                0,
                lambda: self.on_generation_succeeded(
                    content_json,
                    topic,
                    scenario,
                    words,
                    provider,
                    api_key,
                    save_dir,
                    docx_path,
                    mp3_path,
                ),
            )

        except Exception as e:
            self.logger.log(f"\n❌ 發生錯誤: {str(e)}")
            self.root.after(
                0,
                lambda error=str(e): messagebox.showerror(
                    "生成失敗", f"執行過程中發生錯誤：\n{error}"
                ),
            )
        finally:
            self.root.after(0, lambda: self.start_btn.config(state=tk.NORMAL))

    def on_generation_succeeded(
        self,
        content_json,
        topic,
        scenario,
        words,
        provider,
        api_key,
        save_dir,
        docx_path,
        audio_path,
    ):
        self.stop_generated_audio(update_status=False)
        self.generated_docx_path = docx_path
        self.generated_audio_path = audio_path
        self.render_learning_card_preview(content_json)
        self.preview_title_var.set(
            content_json.get("title_zh")
            or content_json.get("title")
            or topic
        )
        self.open_card_btn.config(
            state=tk.NORMAL if os.path.isfile(docx_path) else tk.DISABLED
        )
        self.play_audio_btn.config(
            state=tk.NORMAL if os.path.isfile(audio_path) else tk.DISABLED
        )
        self.pause_audio_btn.config(state=tk.DISABLED)
        self.stop_audio_btn.config(state=tk.DISABLED)
        self.audio_status_var.set(os.path.basename(audio_path))
        self.practice_context = {
            "content": content_json,
            "topic": topic,
            "scenario": scenario,
            "keywords": [
                item["word"].strip()
                for item in content_json.get("words", [])
                if isinstance(item, dict)
                and isinstance(item.get("word"), str)
                and item["word"].strip()
            ] or parse_target_keywords(words),
            "provider": provider,
            "api_key": api_key,
        }
        self.practice_btn.config(state=tk.NORMAL)
        self.tools_menu.entryconfig(self.practice_menu_index, state=tk.NORMAL)
        messagebox.showinfo(
            "成功",
            f"學習卡與語音檔已成功生成！\n\n"
            f"學習卡：{docx_path}\n語音檔：{audio_path}\n\n"
            f"檔案目錄：\n{save_dir}",
        )

    def render_learning_card_preview(self, content):
        preview = self.card_preview_text
        preview.configure(state=tk.NORMAL)
        preview.delete("1.0", tk.END)
        title = content.get("title", "")
        title_zh = content.get("title_zh", "")
        if title:
            preview.insert(tk.END, f"{title}\n", "preview-title")
        if title_zh:
            preview.insert(tk.END, f"{title_zh}\n", "preview-subtitle")

        article = content.get("article", {})
        english_paragraphs = article.get("en", []) if isinstance(article, dict) else []
        chinese_paragraphs = article.get("zh", []) if isinstance(article, dict) else []
        preview.insert(tk.END, "雙語文章\n", "preview-subtitle")
        for index, english in enumerate(english_paragraphs):
            if isinstance(english, str):
                preview.insert(tk.END, f"{english}\n")
            if index < len(chinese_paragraphs) and isinstance(
                chinese_paragraphs[index], str
            ):
                preview.insert(
                    tk.END, f"{chinese_paragraphs[index]}\n", "muted"
                )
            preview.insert(tk.END, "\n")

        preview.insert(tk.END, "目標字詞\n", "preview-subtitle")
        for item in content.get("words", []):
            if not isinstance(item, dict):
                continue
            word = item.get("word", "")
            pronunciation = item.get("kk", "")
            meanings = item.get("pos", [])
            meaning_text = "；".join(
                f"{part.get('pos', '')} {part.get('zh', '')}".strip()
                for part in meanings
                if isinstance(part, dict)
            )
            preview.insert(
                tk.END, f"{word} [{pronunciation}]  {meaning_text}\n", "preview-word"
            )
            usage = item.get("usage")
            if isinstance(usage, str) and usage:
                preview.insert(tk.END, f"用法：{usage}\n")
            forms = item.get("forms")
            if isinstance(forms, str) and forms:
                preview.insert(tk.END, f"詞形變化：{forms}\n")
            for example in item.get("examples", []):
                if isinstance(example, dict):
                    english = example.get("en", "")
                    chinese = example.get("zh", "")
                    if english:
                        preview.insert(tk.END, f"  • {english}\n")
                    if chinese:
                        preview.insert(tk.END, f"    {chinese}\n", "muted")
            for label, values in (
                ("同義字", item.get("synonyms", [])),
                ("常用搭配", item.get("collocations", [])),
            ):
                if isinstance(values, list):
                    values = [value for value in values if isinstance(value, str)]
                    if values:
                        preview.insert(
                            tk.END, f"{label}：{', '.join(values)}\n"
                        )
            listening_items = item.get("listening", [])
            if not isinstance(listening_items, list):
                listening_items = []
            for listening in listening_items:
                if not isinstance(listening, dict):
                    continue
                english = listening.get("en", "")
                chinese = listening.get("zh", "")
                if english:
                    preview.insert(tk.END, f"聽力練習：{english}\n")
                if chinese:
                    preview.insert(tk.END, f"  {chinese}\n", "muted")
        preview.configure(state=tk.DISABLED)
        preview.yview_moveto(0)

    def import_historical_card(self):
        card_path = filedialog.askopenfilename(
            title="匯入歷史 Word 學習卡",
            initialdir=self.output_dir_var.get().strip() or os.getcwd(),
            filetypes=(("Word 學習卡", "*.docx"), ("所有檔案", "*.*")),
            parent=self.root,
        )
        if not card_path:
            return

        try:
            document = Document(card_path)
            lines = []
            for element in document._element.body.iterchildren():
                if element.tag == qn("w:p"):
                    paragraph_text = Paragraph(element, document).text.strip()
                    if paragraph_text:
                        lines.append(paragraph_text)
                elif element.tag == qn("w:tbl"):
                    table = Table(element, document)
                    for row in table.rows:
                        cells = [
                            cell.text.strip().replace("\n", " / ")
                            for cell in row.cells
                        ]
                        if any(cells):
                            lines.append("    ".join(cells))
                    lines.append("")
            preview_text = "\n".join(lines).strip()
            if not preview_text:
                raise ValueError("這份 Word 檔沒有可預覽的文字內容。")
        except (
            OSError,
            ValueError,
            TypeError,
            KeyError,
            zipfile.BadZipFile,
            PackageNotFoundError,
        ) as error:
            messagebox.showerror(
                "匯入學習卡失敗",
                f"無法讀取這份 Word 學習卡：\n{error}",
                parent=self.root,
            )
            return

        self.stop_generated_audio(update_status=False)
        self.generated_docx_path = card_path
        self.open_card_btn.config(state=tk.NORMAL)
        self.preview_title_var.set(os.path.basename(card_path))
        preview = self.card_preview_text
        preview.configure(state=tk.NORMAL)
        preview.delete("1.0", tk.END)
        preview.insert(tk.END, preview_text)
        preview.configure(state=tk.DISABLED)
        preview.yview_moveto(0)

        base_path = card_path[:-len("_Flashcard.docx")] if card_path.lower().endswith(
            "_flashcard.docx"
        ) else os.path.splitext(card_path)[0]
        audio_path = next(
            (
                candidate
                for candidate in (
                    f"{base_path}_article.mp3",
                    f"{base_path}_article.wav",
                )
                if os.path.isfile(candidate)
            ),
            None,
        )
        if audio_path:
            self.set_generated_audio(audio_path)
        else:
            self.generated_audio_path = None
            self.play_audio_btn.config(state=tk.DISABLED)
            self.pause_audio_btn.config(state=tk.DISABLED)
            self.stop_audio_btn.config(state=tk.DISABLED)
            self.audio_status_var.set("未找到配對語音，可按「匯入歷史語音」選擇")

    def import_historical_audio(self):
        audio_path = filedialog.askopenfilename(
            title="匯入歷史語音檔",
            initialdir=(
                os.path.dirname(self.generated_docx_path)
                if self.generated_docx_path
                else self.output_dir_var.get().strip() or os.getcwd()
            ),
            filetypes=(
                ("語音檔", "*.mp3 *.wav"),
                ("MP3 音訊", "*.mp3"),
                ("WAV 音訊", "*.wav"),
                ("所有檔案", "*.*"),
            ),
            parent=self.root,
        )
        if audio_path:
            self.set_generated_audio(audio_path)

    def set_generated_audio(self, audio_path):
        self.stop_generated_audio(update_status=False)
        self.generated_audio_path = audio_path
        self.play_audio_btn.config(state=tk.NORMAL)
        self.pause_audio_btn.config(state=tk.DISABLED)
        self.stop_audio_btn.config(state=tk.DISABLED)
        self.audio_status_var.set(os.path.basename(audio_path))

    def open_generated_card(self):
        if not self.generated_docx_path or not os.path.isfile(
            self.generated_docx_path
        ):
            messagebox.showerror("找不到學習卡", "學習卡檔案不存在，請重新生成。")
            return
        self.open_path(self.generated_docx_path)

    def send_audio_command(self, command, return_output=False):
        if not self.audio_mci:
            raise RuntimeError("音訊播放裝置尚未初始化。")
        output = ctypes.create_unicode_buffer(256)
        result = self.audio_mci.mciSendStringW(
            command,
            ctypes.cast(output, ctypes.c_wchar_p),
            len(output),
            None,
        )
        if result:
            self.audio_mci.mciGetErrorStringW(
                result,
                ctypes.cast(output, ctypes.c_wchar_p),
                len(output),
            )
            raise RuntimeError(output.value or f"Windows 音訊錯誤 {result}")
        return output.value if return_output else None

    def play_generated_audio(self):
        if sys.platform != "win32":
            messagebox.showerror("無法播放", "目前內建播放控制僅支援 Windows。")
            return
        if not self.generated_audio_path or not os.path.isfile(
            self.generated_audio_path
        ):
            messagebox.showerror("找不到語音檔", "語音檔不存在，請重新生成。")
            return
        try:
            self.stop_generated_audio(update_status=False)
            self.audio_mci = ctypes.WinDLL("winmm")
            self.audio_mci.mciSendStringW.argtypes = (
                ctypes.c_wchar_p,
                ctypes.c_wchar_p,
                ctypes.c_uint,
                ctypes.c_void_p,
            )
            self.audio_mci.mciSendStringW.restype = ctypes.c_uint
            self.audio_mci.mciGetErrorStringW.argtypes = (
                ctypes.c_uint,
                ctypes.c_wchar_p,
                ctypes.c_uint,
            )
            self.audio_mci.mciGetErrorStringW.restype = ctypes.c_int
            self.audio_device_alias = f"learning_{uuid.uuid4().hex}"
            audio_type = (
                "waveaudio"
                if self.generated_audio_path.lower().endswith(".wav")
                else "mpegvideo"
            )
            self.send_audio_command(
                f'open "{self.generated_audio_path}" '
                f"type {audio_type} alias {self.audio_device_alias}"
            )
            self.send_audio_command(f"play {self.audio_device_alias}")
        except (OSError, RuntimeError) as error:
            self.close_generated_audio_device()
            messagebox.showerror("播放失敗", f"無法播放語音檔：\n{error}")
            self.audio_status_var.set("播放失敗")
            return

        self.play_audio_btn.config(state=tk.DISABLED)
        self.pause_audio_btn.config(
            state=tk.NORMAL, text="⏸ 暫停", command=self.pause_generated_audio
        )
        self.stop_audio_btn.config(state=tk.NORMAL)
        self.audio_status_var.set("正在播放...")
        self.audio_poll_id = self.root.after(500, self.poll_generated_audio)

    def pause_generated_audio(self):
        if not self.audio_device_alias:
            return
        try:
            self.send_audio_command(f"pause {self.audio_device_alias}")
        except RuntimeError as error:
            messagebox.showerror("暫停失敗", str(error))
            return
        self.pause_audio_btn.config(
            text="▶ 繼續", command=self.resume_generated_audio
        )
        self.audio_status_var.set("已暫停")

    def resume_generated_audio(self):
        if not self.audio_device_alias:
            return
        try:
            self.send_audio_command(f"resume {self.audio_device_alias}")
        except RuntimeError as error:
            messagebox.showerror("繼續播放失敗", str(error))
            return
        self.pause_audio_btn.config(
            text="⏸ 暫停", command=self.pause_generated_audio
        )
        self.audio_status_var.set("正在播放...")

    def poll_generated_audio(self):
        self.audio_poll_id = None
        if not self.audio_device_alias:
            return
        try:
            mode = self.send_audio_command(
                f"status {self.audio_device_alias} mode", return_output=True
            )
        except RuntimeError as error:
            self.close_generated_audio_device()
            self.audio_status_var.set("播放失敗")
            self.play_audio_btn.config(state=tk.NORMAL)
            self.pause_audio_btn.config(state=tk.DISABLED, text="⏸ 暫停")
            self.stop_audio_btn.config(state=tk.DISABLED)
            messagebox.showerror("播放失敗", str(error))
            return
        if mode == "stopped":
            self.close_generated_audio_device()
            self.audio_status_var.set("播放完成")
            self.play_audio_btn.config(state=tk.NORMAL)
            self.pause_audio_btn.config(
                state=tk.DISABLED, text="⏸ 暫停", command=self.pause_generated_audio
            )
            self.stop_audio_btn.config(state=tk.DISABLED)
            return
        self.audio_poll_id = self.root.after(500, self.poll_generated_audio)

    def close_generated_audio_device(self):
        if self.audio_poll_id:
            self.root.after_cancel(self.audio_poll_id)
            self.audio_poll_id = None
        if self.audio_device_alias and self.audio_mci:
            try:
                self.send_audio_command(f"close {self.audio_device_alias}")
            except RuntimeError as error:
                self.logger.log(f"關閉語音播放裝置失敗：{error}")
        self.audio_device_alias = None
        self.audio_mci = None

    def stop_generated_audio(self, update_status=True):
        if self.audio_device_alias and self.audio_mci:
            try:
                self.send_audio_command(f"stop {self.audio_device_alias}")
            except RuntimeError as error:
                self.logger.log(f"停止語音播放失敗：{error}")
            self.close_generated_audio_device()
        elif self.audio_poll_id:
            self.root.after_cancel(self.audio_poll_id)
            self.audio_poll_id = None
        if hasattr(self, "play_audio_btn"):
            self.play_audio_btn.config(
                state=(
                    tk.NORMAL
                    if self.generated_audio_path
                    and os.path.isfile(self.generated_audio_path)
                    else tk.DISABLED
                )
            )
            self.pause_audio_btn.config(
                state=tk.DISABLED, text="⏸ 暫停", command=self.pause_generated_audio
            )
            self.stop_audio_btn.config(state=tk.DISABLED)
        if update_status and hasattr(self, "audio_status_var"):
            self.audio_status_var.set("已停止")

    def close_application(self):
        self.stop_generated_audio(update_status=False)
        self.root.destroy()

    def load_chat_history(self):
        if not os.path.exists(CHAT_HISTORY_FILE):
            return []
        with open(CHAT_HISTORY_FILE, "r", encoding="utf-8") as history_file:
            data = json.load(history_file)
        if not isinstance(data, dict) or not isinstance(data.get("conversations"), list):
            raise ValueError("對話紀錄檔格式不正確，請先備份後再修復 chat_history.json。")
        seen_ids = set()
        for record in data["conversations"]:
            if not isinstance(record, dict):
                raise ValueError("對話紀錄含有無效項目，未載入任何紀錄。")
            if not all(
                isinstance(record.get(field), str)
                for field in ("id", "title", "created_at", "updated_at")
            ):
                raise ValueError("對話紀錄缺少識別碼或時間資訊。")
            if not record["id"] or record["id"] in seen_ids:
                raise ValueError("對話紀錄含有空白或重複的識別碼。")
            seen_ids.add(record["id"])
            context = record.get("context")
            if not isinstance(context, dict) or not isinstance(
                record.get("messages"), list
            ):
                raise ValueError("對話紀錄缺少學習情境或訊息清單。")
            if not all(
                isinstance(context.get(field), str)
                for field in ("title", "topic", "scenario", "provider")
            ) or not isinstance(context.get("keywords"), list) or not all(
                isinstance(keyword, str) for keyword in context["keywords"]
            ) or not isinstance(context.get("words"), list):
                raise ValueError("對話紀錄的學習情境資料格式不正確。")
            for word in context["words"]:
                if (
                    not isinstance(word, dict)
                    or not isinstance(word.get("word"), str)
                    or not isinstance(word.get("usage"), str)
                    or not isinstance(word.get("pos"), list)
                    or not all(
                        isinstance(part, dict) and isinstance(part.get("zh"), str)
                        for part in word["pos"]
                    )
                ):
                    raise ValueError("對話紀錄包含格式不正確的目標單字。")
            for message in record["messages"]:
                if (
                    not isinstance(message, dict)
                    or message.get("role") not in ("user", "assistant")
                    or not isinstance(message.get("content"), str)
                ):
                    raise ValueError("對話紀錄包含無效訊息，未載入任何紀錄。")
        return data["conversations"]

    def save_chat_history(self):
        if self.chat_history_error:
            raise RuntimeError(
                "既有對話紀錄檔無法讀取，為避免覆蓋資料，已停止儲存。"
                f"\n請先備份並修復 {os.path.abspath(CHAT_HISTORY_FILE)}。"
            )
        temp_path = None
        try:
            with tempfile.NamedTemporaryFile(
                mode="w",
                encoding="utf-8",
                dir=os.path.dirname(os.path.abspath(CHAT_HISTORY_FILE)),
                prefix="chat_history_",
                suffix=".tmp",
                delete=False,
            ) as temp_file:
                temp_path = temp_file.name
                json.dump(
                    {"conversations": self.chat_history},
                    temp_file,
                    ensure_ascii=False,
                    indent=2,
                )
            os.replace(temp_path, CHAT_HISTORY_FILE)
            self.chat_history_error = None
        except (OSError, TypeError, ValueError) as error:
            if temp_path and os.path.exists(temp_path):
                try:
                    os.remove(temp_path)
                except OSError as cleanup_error:
                    self.logger.log(f"清理暫存對話紀錄失敗：{cleanup_error}")
            raise RuntimeError(f"儲存對話紀錄失敗：{error}") from error

    def create_chat_record(self):
        content = self.practice_context["content"]
        target_words = []
        for item in content.get("words", []):
            if not isinstance(item, dict):
                continue
            parts = item.get("pos", [])
            if not isinstance(parts, list):
                parts = []
            target_words.append(
                {
                    "word": (
                        item.get("word", "")
                        if isinstance(item.get("word", ""), str)
                        else ""
                    ),
                    "pos": [
                        {
                            "zh": (
                                part.get("zh", "")
                                if isinstance(part.get("zh", ""), str)
                                else ""
                            )
                        }
                        for part in parts
                        if isinstance(part, dict)
                    ],
                    "usage": (
                        item.get("usage", "")
                        if isinstance(item.get("usage", ""), str)
                        else ""
                    ),
                }
            )
        now = datetime.now().astimezone().isoformat(timespec="seconds")
        title = content.get("title_zh") or content.get("title")
        if not isinstance(title, str) or not title.strip():
            title = self.practice_context["topic"]
        return {
            "id": str(uuid.uuid4()),
            "title": title,
            "created_at": now,
            "updated_at": now,
            "context": {
                "title": (
                    content.get("title", "")
                    if isinstance(content.get("title", ""), str)
                    else ""
                ),
                "topic": self.practice_context["topic"],
                "scenario": self.practice_context["scenario"],
                "keywords": self.practice_context["keywords"],
                "words": target_words,
                "provider": self.practice_context["provider"],
            },
            "messages": [],
        }

    def persist_current_conversation(self):
        if not self.practice_history_id and not self.practice_messages:
            return
        record = next(
            (
                item
                for item in self.chat_history
                if item["id"] == self.practice_history_id
            ),
            None,
        )
        if not self.practice_messages:
            if record is None:
                return
            self.chat_history = [
                item
                for item in self.chat_history
                if item["id"] != self.practice_history_id
            ]
            try:
                self.save_chat_history()
            except RuntimeError as error:
                self.append_practice_text("系統提示", str(error), error=True)
            if hasattr(self, "history_tree") and self.history_tree.winfo_exists():
                self.refresh_conversation_history()
            return
        if not self.practice_history_id:
            self.practice_history_id = str(uuid.uuid4())
            record = None
        if record is None:
            record = self.create_chat_record()
            record["id"] = self.practice_history_id
            self.chat_history.append(record)
        record["messages"] = [dict(message) for message in self.practice_messages]
        record["updated_at"] = datetime.now().astimezone().isoformat(timespec="seconds")
        try:
            self.save_chat_history()
        except RuntimeError as error:
            self.append_practice_text("系統提示", str(error), error=True)
            return
        if hasattr(self, "history_tree") and self.history_tree.winfo_exists():
            self.refresh_conversation_history()

    def open_conversation_history(self):
        if self.chat_history_error:
            messagebox.showerror(
                "無法讀取對話紀錄",
                f"{self.chat_history_error}\n\n"
                f"請先備份並修復：{os.path.abspath(CHAT_HISTORY_FILE)}",
                parent=self.root,
            )
            return
        if hasattr(self, "history_window") and self.history_window.winfo_exists():
            self.history_window.lift()
            self.refresh_conversation_history()
            return

        window = tk.Toplevel(self.root)
        self.history_window = window
        window.title("AI 對話紀錄")
        window.geometry("760x480")
        window.minsize(600, 360)
        window.transient(self.root)
        window.columnconfigure(0, weight=1)
        window.rowconfigure(1, weight=1)

        history_header = tk.Frame(
            window, bg="#E9F5F4", padx=14, pady=10
        )
        history_header.grid(row=0, column=0, sticky=tk.EW)
        tk.Label(
            history_header,
            text="對話會自動儲存在本機。選取一筆紀錄即可接續練習。",
            bg="#E9F5F4",
            fg=APP_COLORS["navy"],
            font=("Microsoft JhengHei", 10, "bold"),
        ).pack(anchor=tk.W)
        list_frame = ttk.Frame(window, padding=(12, 0, 12, 8))
        list_frame.grid(row=1, column=0, sticky=tk.NSEW)
        list_frame.columnconfigure(0, weight=1)
        list_frame.rowconfigure(0, weight=1)
        self.history_tree = ttk.Treeview(
            list_frame,
            columns=("scenario", "messages", "updated"),
            show="tree headings",
            selectmode="browse",
            style="History.Treeview",
        )
        self.history_tree.heading("#0", text="學習主題")
        self.history_tree.heading("scenario", text="情境")
        self.history_tree.heading("messages", text="對話數")
        self.history_tree.heading("updated", text="最後更新")
        self.history_tree.column("#0", width=300, anchor=tk.W)
        self.history_tree.column("scenario", width=130, anchor=tk.W)
        self.history_tree.column("messages", width=80, anchor=tk.CENTER)
        self.history_tree.column("updated", width=180, anchor=tk.W)
        history_scrollbar = ttk.Scrollbar(
            list_frame, command=self.history_tree.yview
        )
        self.history_tree.configure(yscrollcommand=history_scrollbar.set)
        self.history_tree.grid(row=0, column=0, sticky=tk.NSEW)
        history_scrollbar.grid(row=0, column=1, sticky=tk.NS)
        self.history_tree.bind("<Double-1>", self.continue_selected_conversation)

        buttons = ttk.Frame(window, padding=(12, 0, 12, 12))
        buttons.grid(row=2, column=0, sticky=tk.EW)
        ttk.Button(
            buttons,
            text="開始新對話",
            command=self.start_new_practice_from_history,
            style="Secondary.TButton",
        ).pack(side=tk.LEFT)
        ttk.Button(
            buttons,
            text="刪除選取紀錄",
            command=self.delete_selected_conversation,
        ).pack(side=tk.LEFT, padx=8)
        ttk.Button(
            buttons,
            text="接續選取對話",
            command=self.continue_selected_conversation,
            style="Accent.TButton",
        ).pack(side=tk.RIGHT)
        self.refresh_conversation_history()

    def refresh_conversation_history(self):
        if not hasattr(self, "history_tree") or not self.history_tree.winfo_exists():
            return
        existing_items = self.history_tree.get_children()
        if existing_items:
            self.history_tree.delete(*existing_items)
        records = sorted(
            self.chat_history, key=lambda item: item["updated_at"], reverse=True
        )
        for record in records:
            updated = record["updated_at"]
            try:
                updated = datetime.fromisoformat(updated).astimezone().strftime(
                    "%Y-%m-%d %H:%M"
                )
            except ValueError:
                pass
            self.history_tree.insert(
                "",
                tk.END,
                iid=record["id"],
                text=record["title"],
                values=(
                    record["context"].get("scenario", ""),
                    len(record["messages"]),
                    updated,
                ),
            )

    def selected_conversation_record(self):
        selected = self.history_tree.selection()
        if not selected:
            messagebox.showinfo(
                "請選擇對話", "請先在清單中選擇一筆對話紀錄。",
                parent=self.history_window,
            )
            return None
        return next(
            (record for record in self.chat_history if record["id"] == selected[0]),
            None,
        )

    def continue_selected_conversation(self, event=None):
        if event is not None and event.widget is not self.history_tree:
            return
        record = self.selected_conversation_record()
        if record:
            self.open_practice_window(history_id=record["id"])

    def start_new_practice_from_history(self):
        if not self.practice_context:
            messagebox.showwarning(
                "尚無學習卡",
                "請先生成一份學習卡，再開始新的對話練習。",
                parent=self.history_window,
            )
            return
        self.open_practice_window()

    def delete_selected_conversation(self):
        record = self.selected_conversation_record()
        if not record:
            return
        if not messagebox.askyesno(
            "確認刪除",
            f"確定刪除「{record['title']}」的對話紀錄嗎？此動作無法復原。",
            parent=self.history_window,
        ):
            return
        self.chat_history = [
            item for item in self.chat_history if item["id"] != record["id"]
        ]
        try:
            self.save_chat_history()
        except RuntimeError as error:
            self.chat_history.append(record)
            self.refresh_conversation_history()
            messagebox.showerror(
                "刪除失敗", str(error), parent=self.history_window
            )
            return
        if self.practice_history_id == record["id"]:
            self.practice_history_id = None
            self.close_practice_window()
        self.refresh_conversation_history()

    def restore_practice_context(self, record):
        context = record["context"]
        provider = context.get("provider", "Gemini")
        if provider not in AI_PROVIDERS:
            raise ValueError(f"紀錄中的 AI 服務不支援：{provider}")
        api_key = self.unsaved_api_keys.get(provider) or self.get_provider_api_key(provider)
        if provider == self.provider_combo.get() and self.api_key_entry.get().strip():
            api_key = self.api_key_entry.get().strip()
        if not api_key:
            messagebox.showwarning(
                "缺少 API Key",
                f"此對話使用 {provider}。請先設定有效的 API Key，再接續對話。",
                parent=self.history_window if hasattr(self, "history_window") else self.root,
            )
            return False
        self.practice_context = {
            "content": {
                "title": context.get("title", ""),
                "words": context.get("words", []),
            },
            "topic": context.get("topic", ""),
            "scenario": context.get("scenario", ""),
            "keywords": context.get("keywords", []),
            "provider": provider,
            "api_key": api_key,
        }
        return True

    def open_practice_window(self, history_id=None):
        history_record = None
        if history_id:
            history_record = next(
                (item for item in self.chat_history if item["id"] == history_id),
                None,
            )
            if not history_record:
                messagebox.showerror("找不到對話", "這筆對話紀錄已不存在。")
                return
            if not self.restore_practice_context(history_record):
                return
        elif not self.practice_context:
            messagebox.showwarning("提示", "請先成功生成一份學習卡，再開始對話練習。")
            return
        if self.practice_window and self.practice_window.winfo_exists():
            self.close_practice_window()

        self.practice_request_token += 1
        self.practice_history_id = (
            history_record["id"] if history_record else str(uuid.uuid4())
        )
        self.practice_messages = (
            [dict(message) for message in history_record["messages"]]
            if history_record
            else []
        )
        window = tk.Toplevel(self.root)
        self.practice_window = window
        window.title(
            f"AI 英文對話練習 - {history_record['title']}"
            if history_record
            else "AI 英文對話練習"
        )
        window.geometry("700x650")
        window.minsize(560, 480)
        window.transient(self.root)
        window.protocol("WM_DELETE_WINDOW", self.close_practice_window)

        practice_header = tk.Frame(
            window, bg=APP_COLORS["navy"], padx=14, pady=10
        )
        practice_header.pack(fill=tk.X)
        tk.Label(
            practice_header,
            text=(
                f"主題：{self.practice_context['topic']}　"
                f"情境：{self.practice_context['scenario']}"
            ),
            bg=APP_COLORS["navy"],
            fg="#FFFFFF",
            font=("Microsoft JhengHei", 10, "bold"),
        ).pack(anchor=tk.W)

        conversation_frame = ttk.Frame(window, padding=(12, 0))
        conversation_frame.pack(fill=tk.BOTH, expand=True)
        self.practice_text = tk.Text(
            conversation_frame,
            wrap=tk.WORD,
            state=tk.DISABLED,
            font=("Microsoft JhengHei", 10),
            bg="#FBFDFF",
            fg=APP_COLORS["text"],
            relief=tk.FLAT,
            padx=10,
            pady=8,
        )
        scrollbar = ttk.Scrollbar(
            conversation_frame, command=self.practice_text.yview
        )
        self.practice_text.configure(yscrollcommand=scrollbar.set)
        self.practice_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        self.practice_text.tag_configure("speaker", font=("Microsoft JhengHei", 10, "bold"))
        self.practice_text.tag_configure(
            "speaker-ai",
            foreground=APP_COLORS["blue"],
            font=("Microsoft JhengHei", 10, "bold"),
        )
        self.practice_text.tag_configure(
            "speaker-user",
            foreground=APP_COLORS["teal"],
            font=("Microsoft JhengHei", 10, "bold"),
        )
        self.practice_text.tag_configure("error", foreground="#B03A2E")
        if self.practice_messages:
            for message in self.practice_messages:
                speaker = "你" if message["role"] == "user" else "AI"
                self.append_practice_text(speaker, message["content"])
        else:
            self.append_practice_text(
                "AI",
                "Hi! Let's practice together. Choose a prompt below or type your own response.",
            )

        hints_frame = ttk.LabelFrame(
            window,
            text="英文提示句（點選後可編輯）",
            padding=8,
            style="Card.TLabelframe",
        )
        hints_frame.pack(fill=tk.X, padx=12, pady=(10, 6))
        self.practice_input = tk.StringVar()
        for index, hint in enumerate(
            (
                "Could you tell me more?",
                "I think that...",
                "Could you repeat that, please?",
                "Let me check...",
            )
        ):
            ttk.Button(
                hints_frame,
                text=hint,
                command=lambda phrase=hint: self.insert_practice_hint(phrase),
                style="Secondary.TButton",
            ).grid(row=index // 2, column=index % 2, sticky=tk.W, padx=3, pady=2)

        input_frame = ttk.Frame(window, padding=(12, 4, 12, 12))
        input_frame.pack(fill=tk.X)
        self.practice_entry = ttk.Entry(
            input_frame, textvariable=self.practice_input, font=("Microsoft JhengHei", 10)
        )
        self.practice_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 8))
        self.practice_entry.bind("<Return>", self.send_practice_message)
        self.practice_send_btn = ttk.Button(
            input_frame,
            text="送出",
            command=self.send_practice_message,
            style="Accent.TButton",
        )
        self.practice_send_btn.pack(side=tk.RIGHT)
        self.practice_speech_btn = ttk.Button(
            input_frame,
            text="🎙 語音",
            command=lambda: self.start_speech_input(
                target_entry=self.practice_entry,
                button=self.practice_speech_btn,
                language="en-US",
            ),
            style="Secondary.TButton",
        )
        self.practice_speech_btn.pack(side=tk.RIGHT, padx=(0, 8))
        self.practice_entry.focus_set()

    def close_practice_window(self):
        self.practice_request_token += 1
        if self.practice_window and self.practice_window.winfo_exists():
            self.practice_window.destroy()
        self.practice_window = None

    def append_practice_text(self, speaker, text, error=False):
        self.practice_text.configure(state=tk.NORMAL)
        speaker_tag = (
            "speaker-ai" if speaker == "AI" else "speaker-user"
        )
        self.practice_text.insert(tk.END, f"{speaker}：", speaker_tag)
        self.practice_text.insert(
            tk.END, f"{text}\n\n", "error" if error else ()
        )
        self.practice_text.configure(state=tk.DISABLED)
        self.practice_text.see(tk.END)

    def insert_practice_hint(self, phrase):
        current = self.practice_input.get()
        self.practice_input.set(f"{current} {phrase}".strip())
        self.practice_entry.focus_set()
        self.practice_entry.icursor(tk.END)

    def build_practice_system_prompt(self):
        context = self.practice_context
        content = context["content"]
        words = [
            {
                "word": item.get("word", ""),
                "meaning": " ".join(
                    part.get("zh", "")
                    for part in item.get("pos", [])
                    if isinstance(part, dict)
                ),
                "usage": item.get("usage", ""),
            }
            for item in content.get("words", [])
            if isinstance(item, dict)
        ]
        learning_context = json.dumps(
            {
                "title": content.get("title", ""),
                "topic": context["topic"],
                "scenario": context["scenario"],
                "target_words": words,
            },
            ensure_ascii=False,
        )
        return (
            "You are a friendly English conversation tutor for a Taiwanese learner. "
            "Practice the situation and vocabulary in this learning-card context: "
            f"{learning_context}\n"
            "Speak mainly in natural, concise English suitable for an intermediate beginner. "
            "Keep the role-play moving by asking one short follow-up question at a time. "
            "When the learner makes an English mistake, gently show a brief correction "
            "and explanation in Traditional Chinese, then continue the conversation in English. "
            "Encourage the learner to use the target words, but do not force them unnaturally."
        )

    def send_practice_message(self, event=None):
        message = self.practice_input.get().strip()
        if not message:
            return "break" if event else None
        if not self.practice_window or not self.practice_window.winfo_exists():
            return "break" if event else None

        self.practice_input.set("")
        self.append_practice_text("你", message)
        self.practice_messages.append({"role": "user", "content": message})
        self.persist_current_conversation()
        self.practice_send_btn.config(state=tk.DISABLED)
        self.practice_entry.config(state=tk.DISABLED)
        token = self.practice_request_token
        window = self.practice_window
        messages = [dict(item) for item in self.practice_messages]
        context = self.practice_context
        system_prompt = self.build_practice_system_prompt()

        threading.Thread(
            target=self.run_practice_turn,
            args=(token, window, context, system_prompt, messages),
            daemon=True,
        ).start()
        return "break" if event else None

    def run_practice_turn(self, token, window, context, system_prompt, messages):
        try:
            reply = self.call_chat_api(
                context["provider"],
                context["api_key"],
                system_prompt,
                messages,
            )
            self.root.after(
                0, lambda: self.finish_practice_turn(token, window, reply, None)
            )
        except Exception as error:
            self.root.after(
                0,
                lambda message=str(error): self.finish_practice_turn(
                    token, window, None, message
                ),
            )

    def finish_practice_turn(self, token, window, reply, error):
        if (
            token != self.practice_request_token
            or window is not self.practice_window
            or not window.winfo_exists()
        ):
            return

        if error:
            self.practice_messages.pop()
            self.append_practice_text("錯誤", error, error=True)
            self.persist_current_conversation()
        else:
            self.practice_messages.append({"role": "assistant", "content": reply})
            self.append_practice_text("AI", reply)
            self.persist_current_conversation()
            speech_text = re.sub(r"[\u3400-\u9fff]+", " ", reply)
            speech_text = re.sub(r"\s+", " ", speech_text).strip()
            if speech_text:
                voice_code = self.voice_var.get().split()[0]
                self.practice_send_btn.config(
                    state=tk.DISABLED, text="語音回覆中..."
                )
                self.practice_entry.config(state=tk.DISABLED)
                threading.Thread(
                    target=self.run_practice_speech,
                    args=(token, window, speech_text, voice_code),
                    daemon=True,
                ).start()
                return
        self.practice_send_btn.config(state=tk.NORMAL)
        self.practice_entry.config(state=tk.NORMAL)
        self.practice_entry.focus_set()

    def run_practice_speech(self, token, window, text, voice_code):
        audio_path = None
        try:
            with tempfile.NamedTemporaryFile(suffix=".mp3", delete=False) as audio_file:
                audio_path = audio_file.name

            async def synthesize():
                communicate = edge_tts.Communicate(text, voice_code)
                await communicate.save(audio_path)

            asyncio.run(synthesize())
            if os.path.getsize(audio_path) == 0:
                raise RuntimeError("Edge TTS 沒有產生有效的語音回覆。")
            self.play_practice_audio(audio_path)
            error = None
        except (
            edge_tts.exceptions.EdgeTTSException,
            aiohttp.ClientError,
            asyncio.TimeoutError,
            OSError,
            RuntimeError,
            ctypes.ArgumentError,
        ) as speech_error:
            error = f"語音回覆失敗：{speech_error}"
        finally:
            if audio_path and os.path.exists(audio_path):
                try:
                    os.remove(audio_path)
                except OSError as cleanup_error:
                    self.logger.log(f"清理語音回覆暫存檔失敗：{cleanup_error}")
        self.root.after(
            0,
            lambda: self.finish_practice_speech(token, window, error),
        )

    @staticmethod
    def play_practice_audio(audio_path):
        if sys.platform != "win32":
            raise RuntimeError("目前的語音回覆播放功能僅支援 Windows。")

        winmm = ctypes.WinDLL("winmm")
        device_alias = f"practice_{uuid.uuid4().hex}"
        command_buffer = ctypes.create_unicode_buffer(256)
        winmm.mciSendStringW.argtypes = (
            ctypes.c_wchar_p,
            ctypes.c_wchar_p,
            ctypes.c_uint,
            ctypes.c_void_p,
        )
        winmm.mciSendStringW.restype = ctypes.c_uint
        winmm.mciGetErrorStringW.argtypes = (
            ctypes.c_uint,
            ctypes.c_wchar_p,
            ctypes.c_uint,
        )
        winmm.mciGetErrorStringW.restype = ctypes.c_int
        opened = False

        def send_mci_command(command):
            result = winmm.mciSendStringW(
                command,
                ctypes.cast(command_buffer, ctypes.c_wchar_p),
                len(command_buffer),
                None,
            )
            if result:
                winmm.mciGetErrorStringW(
                    result,
                    ctypes.cast(command_buffer, ctypes.c_wchar_p),
                    len(command_buffer),
                )
                raise RuntimeError(
                    command_buffer.value or f"Windows 音訊錯誤 {result}"
                )

        try:
            send_mci_command(
                f'open "{audio_path}" type mpegvideo alias {device_alias}'
            )
            opened = True
            send_mci_command(f"play {device_alias} wait")
        finally:
            if opened:
                send_mci_command(f"close {device_alias}")

    def finish_practice_speech(self, token, window, error):
        if (
            token != self.practice_request_token
            or window is not self.practice_window
            or not window.winfo_exists()
        ):
            return
        if error:
            self.append_practice_text("系統提示", error, error=True)
        self.practice_send_btn.config(state=tk.NORMAL, text="送出")
        self.practice_entry.config(state=tk.NORMAL)
        self.practice_entry.focus_set()

    def call_chat_api(self, provider, api_key, system_prompt, messages):
        if provider == "OpenAI":
            response = requests.post(
                "https://api.openai.com/v1/chat/completions",
                headers={"Authorization": f"Bearer {api_key}"},
                json={
                    "model": "gpt-4o-mini",
                    "messages": [{"role": "system", "content": system_prompt}, *messages],
                    "temperature": 0.6,
                },
                timeout=90,
            )
            response_data = self.parse_chat_response(response, provider)
            try:
                return response_data["choices"][0]["message"]["content"].strip()
            except (KeyError, IndexError, TypeError, AttributeError) as error:
                raise RuntimeError(f"{provider} 回傳內容格式不正確：{error}") from error

        if provider == "Claude":
            response = requests.post(
                "https://api.anthropic.com/v1/messages",
                headers={
                    "x-api-key": api_key,
                    "anthropic-version": "2023-06-01",
                    "Content-Type": "application/json",
                },
                json={
                    "model": "claude-haiku-5-5",
                    "max_tokens": 1024,
                    "system": system_prompt,
                    "messages": messages,
                },
                timeout=90,
            )
            response_data = self.parse_chat_response(response, provider)
            try:
                return "".join(
                    block["text"]
                    for block in response_data["content"]
                    if block.get("type") == "text"
                ).strip()
            except (KeyError, TypeError) as error:
                raise RuntimeError(f"{provider} 回傳內容格式不正確：{error}") from error

        if provider == "Gemini":
            contents = [
                {
                    "role": "model" if item["role"] == "assistant" else "user",
                    "parts": [{"text": item["content"]}],
                }
                for item in messages
            ]
            response = requests.post(
                "https://generativelanguage.googleapis.com/v1beta/models/"
                f"{GEMINI_MODEL}:generateContent",
                params={"key": api_key},
                headers={"Content-Type": "application/json"},
                json={
                    "systemInstruction": {"parts": [{"text": system_prompt}]},
                    "contents": contents,
                    "generationConfig": {"temperature": 0.6},
                },
                timeout=90,
            )
            response_data = self.parse_chat_response(response, provider)
            try:
                return "".join(
                    part["text"]
                    for part in response_data["candidates"][0]["content"]["parts"]
                    if "text" in part
                ).strip()
            except (KeyError, IndexError, TypeError) as error:
                raise RuntimeError(f"{provider} 回傳內容格式不正確：{error}") from error

        raise ValueError(f"不支援的 AI 服務：{provider}")

    @staticmethod
    def parse_chat_response(response, provider):
        try:
            response.raise_for_status()
        except requests.RequestException as error:
            try:
                error_message = response.json().get("error", {}).get("message")
            except (ValueError, AttributeError):
                error_message = None
            detail = error_message or response.text or str(error)
            raise RuntimeError(f"{provider} 對話請求失敗：{detail}") from error
        try:
            return response.json()
        except ValueError as error:
            raise RuntimeError(f"{provider} 回傳的內容不是有效 JSON。") from error

    def call_ai_api(self, provider, api_key, topic, words, scenario):
        keywords = parse_target_keywords(words)
        prompt = build_learning_prompt(topic, keywords, scenario)
        if provider == "Gemini":
            return self.call_gemini_api(api_key, topic, words, scenario)

        if provider == "OpenAI":
            url = "https://api.openai.com/v1/chat/completions"
            headers = {"Authorization": f"Bearer {api_key}"}
            payload = {
                "model": "gpt-4o-mini",
                "messages": [{"role": "user", "content": prompt}],
                "response_format": {"type": "json_object"},
                "temperature": 0.2,
            }
        elif provider == "Claude":
            url = "https://api.anthropic.com/v1/messages"
            headers = {
                "x-api-key": api_key,
                "anthropic-version": "2023-06-01",
                "Content-Type": "application/json",
            }
            payload = {
                "model": "claude-haiku-5-5",
                "max_tokens": 8192,
                "system": "請遵守使用者指示，只輸出有效 JSON，不要使用 Markdown 程式碼區塊。",
                "messages": [{"role": "user", "content": prompt}],
            }
        else:
            raise ValueError(f"不支援的 AI 服務：{provider}")

        try:
            response = requests.post(
                url, headers=headers, json=payload, timeout=120
            )
        except requests.RequestException as error:
            raise RuntimeError(f"{provider} 連線異常：{error}") from error
        if response.status_code != 200:
            try:
                body = response.json()
                error_message = (
                    body.get("error", {}).get("message")
                    or body.get("message")
                    or response.text
                )
            except (ValueError, AttributeError):
                error_message = response.text
            raise RuntimeError(f"{provider} HTTP {response.status_code}: {error_message}")

        try:
            response_data = response.json()
            if provider == "OpenAI":
                text = response_data["choices"][0]["message"]["content"]
            else:
                text = "".join(
                    block["text"] for block in response_data["content"]
                    if block.get("type") == "text"
                )
        except (KeyError, IndexError, TypeError) as error:
            raise RuntimeError(f"{provider} 回傳內容格式不正確：{error}") from error
        return parse_ai_json(text, provider)

    def call_gemini_api(self, api_key, topic, words, scenario="商業"):
        keywords = parse_target_keywords(words)
        candidate_models = [GEMINI_MODEL]
        prompt = build_learning_prompt(topic, keywords, scenario)

        last_error = None
        url_template = "https://generativelanguage.googleapis.com/v1beta/models/{}:generateContent"
        headers = {"Content-Type": "application/json"}
        payload = {
            "contents": [{"parts": [{"text": prompt}]}],
            "generationConfig": {
                "responseMimeType": "application/json",
                "temperature": 0.2
            }
        }

        for model in candidate_models:
            url = url_template.format(model)
            self.logger.log(f"   嘗試使用 Gemini 模型: {model}...")

            for attempt in range(3):
                try:
                    response = requests.post(
                        url,
                        params={"key": api_key},
                        headers=headers,
                        json=payload,
                        timeout=60
                    )
                except requests.RequestException as e:
                    last_error = f"連線異常: {e}"
                    self.logger.log(f"   ⚠️ {last_error}")
                    if attempt < 2:
                        time.sleep(2 ** attempt)
                    continue

                if response.status_code == 200:
                    try:
                        res_data = response.json()
                        parts = res_data["candidates"][0]["content"]["parts"]
                        text_content = "".join(
                            part["text"] for part in parts if "text" in part
                        ).strip()
                        return parse_ai_json(text_content, "Gemini")
                    except (KeyError, IndexError, TypeError, ValueError) as e:
                        raise RuntimeError(
                            f"Gemini 已回應，但回傳內容不是預期的 JSON 格式: {e}"
                        ) from e

                try:
                    error_data = response.json()
                    error_message = error_data.get("error", {}).get("message", response.text)
                except (ValueError, AttributeError):
                    error_message = response.text
                last_error = f"HTTP {response.status_code}: {error_message}"

                if response.status_code == 404:
                    self.logger.log(f"   ⚠️ 模型 {model} 不可用，切換備用模型...")
                    break

                if response.status_code in (429, 500, 502, 503, 504):
                    if attempt < 2:
                        wait_time = 2 ** attempt + 1
                        self.logger.log(
                            f"   ⏳ Gemini 暫時無法處理請求 ({response.status_code})，"
                            f"{wait_time} 秒後重試..."
                        )
                        time.sleep(wait_time)
                        continue
                    break

                if response.status_code in (400, 401, 403):
                    raise RuntimeError(
                        f"{last_error}\n請確認使用的是 Google AI Studio 的 Gemini API Key，"
                        "且該金鑰已啟用 Gemini API。"
                    )

                raise RuntimeError(last_error)

        raise RuntimeError(
            f"所有 Gemini 模型呼叫失敗。最後錯誤: {last_error or '沒有收到 API 回應'}"
        )

    def generate_docx(self, data, save_dir, topic, keywords):
        doc = Document()
        keyword_pattern = build_keyword_pattern(keywords)

        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)
            footer = section.footer
            page_number = footer.paragraphs[0]
            page_number.alignment = WD_ALIGN_PARAGRAPH.CENTER
            page_number.add_run("第 ")
            field_run = page_number.add_run()
            field_begin = OxmlElement("w:fldChar")
            field_begin.set(qn("w:fldCharType"), "begin")
            field_run._r.append(field_begin)
            instruction_run = page_number.add_run()
            page_instruction = OxmlElement("w:instrText")
            page_instruction.set(qn("xml:space"), "preserve")
            page_instruction.text = " PAGE "
            instruction_run._r.append(page_instruction)
            field_separator_run = page_number.add_run()
            field_separator = OxmlElement("w:fldChar")
            field_separator.set(qn("w:fldCharType"), "separate")
            field_separator_run._r.append(field_separator)
            page_number.add_run("1")
            field_end_run = page_number.add_run()
            field_end = OxmlElement("w:fldChar")
            field_end.set(qn("w:fldCharType"), "end")
            field_end_run._r.append(field_end)
            page_number.add_run(" 頁")

        p_title = doc.add_paragraph()
        for r_en in add_highlighted_text(
            p_title, data['title'] + "\n", keyword_pattern
        ):
            r_en.font.name = "Arial"
            r_en.font.size = Pt(16)
            r_en.font.bold = True

        r_zh = p_title.add_run(data['title_zh'])
        r_zh.font.name = "Microsoft JhengHei"
        r_zh.font.size = Pt(14)
        r_zh.font.color.rgb = RGBColor(100, 100, 100)

        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "English"
        hdr_cells[1].text = "中文對照"

        for cell in hdr_cells:
            shading = parse_xml(r'<w:shd {} w:fill="F0F4F8"/>'.format(nsdecls('w')))
            cell._tc.get_or_add_tcPr().append(shading)

        for en_p, zh_p in zip(data['article']['en'], data['article']['zh']):
            row_cells = table.add_row().cells
            for run in set_cell_highlighted_text(row_cells[0], en_p, keyword_pattern):
                run.font.name = "Arial"
            set_cell_highlighted_text(row_cells[1], zh_p, keyword_pattern)

        doc.add_paragraph()

        for word_info in data['words']:
            kk_text = word_info["kk"].strip().strip("[]")
            p_w = doc.add_paragraph()
            for r_w in add_highlighted_text(
                p_w, f"{word_info['word']} ", keyword_pattern
            ):
                r_w.font.name = "Arial"
                r_w.font.size = Pt(14)
                r_w.font.bold = True

            pronunciation = p_w.add_run(f"[{kk_text}] ")
            pronunciation.font.name = "Arial"
            pronunciation.font.size = Pt(11)
            pronunciation.font.color.rgb = RGBColor(80, 80, 80)

            pos_str = " ".join([f"{p['pos']} {p['zh']}" for p in word_info['pos']])
            for r_pos in add_highlighted_text(p_w, pos_str, keyword_pattern):
                r_pos.font.name = "Microsoft JhengHei"
                r_pos.font.size = Pt(11)

            w_table = doc.add_table(rows=0, cols=2)
            set_table_borders(w_table)

            def add_row(label, val):
                r = w_table.add_row().cells
                r[0].text = label
                set_cell_highlighted_text(r[1], val, keyword_pattern)
                r[0].width = Inches(1.2)
                r[1].width = Inches(5.0)

            add_row("KK 音標", f"[{kk_text}]")
            add_row("詞形變化", word_info.get('forms', ''))
            add_row("用法說明", word_info.get('usage', ''))

            ex_str = "\n".join([f"• {ex['en']}\n  {ex['zh']}" for ex in word_info.get('examples', [])])
            add_row("例句", ex_str)

            add_row("同義字", ", ".join(word_info.get('synonyms', [])))
            add_row("搭配詞", ", ".join(word_info.get('collocations', [])))

            doc.add_paragraph()

        doc.add_page_break()
        notes_heading = doc.add_paragraph()
        notes_heading.paragraph_format.space_after = Pt(12)
        notes_title = notes_heading.add_run("筆記欄")
        notes_title.font.name = "Microsoft JhengHei"
        notes_title.font.size = Pt(16)
        notes_title.font.bold = True

        notes_table = doc.add_table(rows=20, cols=1)
        notes_table.autofit = False
        notes_table.columns[0].width = Inches(6.5)
        for row in notes_table.rows:
            row.height = Inches(0.38)
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            cell = row.cells[0]
            cell.width = Inches(6.5)
            borders = OxmlElement("w:tcBorders")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "dashed")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "0")
            bottom.set(qn("w:color"), "7F8C8D")
            borders.append(bottom)
            cell._tc.get_or_add_tcPr().append(borders)

        filename = f"{topic}_Flashcard.docx"
        file_path = os.path.join(save_dir, filename)
        doc.save(file_path)
        return file_path

    def generate_audio(self, data, save_dir, topic, voice_code):
        article_text = " ".join(data['article']['en'])
        if not article_text.strip():
            raise ValueError("文章沒有英文內容，無法生成語音。")

        mp3_path = os.path.join(save_dir, f"{topic}_article.mp3")
        with tempfile.NamedTemporaryFile(dir=save_dir, suffix=".mp3", delete=False) as temp_file:
            temp_mp3_path = temp_file.name

        try:
            async def _tts():
                communicate = edge_tts.Communicate(article_text, voice_code)
                await communicate.save(temp_mp3_path)

            try:
                asyncio.run(_tts())
                if os.path.getsize(temp_mp3_path) == 0:
                    raise edge_tts.exceptions.NoAudioReceived("Edge TTS returned an empty audio file.")
                os.replace(temp_mp3_path, mp3_path)
                return mp3_path
            except (
                edge_tts.exceptions.EdgeTTSException,
                aiohttp.ClientError,
                asyncio.TimeoutError,
            ) as tts_error:
                self.logger.log(
                    "   ⚠️ Edge TTS 服務無法產生音訊，改用 Windows 內建英文語音輸出 WAV..."
                )
                if os.path.exists(temp_mp3_path):
                    os.remove(temp_mp3_path)
                return self.generate_windows_audio(
                    article_text, save_dir, topic, voice_code, tts_error
                )
        finally:
            if os.path.exists(temp_mp3_path):
                os.remove(temp_mp3_path)

    def generate_windows_audio(
        self, article_text, save_dir, topic, voice_code, tts_error
    ):
        if sys.platform != "win32":
            raise RuntimeError(
                f"Edge TTS 無法產生音訊，且 Windows 離線語音備援無法在此系統使用: {tts_error}"
            ) from tts_error

        wav_path = os.path.join(save_dir, f"{topic}_article.wav")
        with tempfile.NamedTemporaryFile(dir=save_dir, suffix=".wav", delete=False) as temp_file:
            temp_wav_path = temp_file.name

        script = """
$ErrorActionPreference = 'Stop'
$text = [Text.Encoding]::UTF8.GetString([Convert]::FromBase64String($env:TOIC_TTS_TEXT))
$synth = New-Object -ComObject SAPI.SpVoice
$gender = if ($env:TOIC_TTS_VOICE -match '-(Guy|Ryan)Neural$') { 'Male' } else { 'Female' }
$language = if ($env:TOIC_TTS_VOICE -like 'en-GB-*') { '809' } else { '409' }
$englishVoices = @($synth.GetVoices() | Where-Object { $_.GetDescription() -match 'English' })
$voice = $englishVoices | Where-Object {
    $_.GetAttribute('Gender') -eq $gender -and
    $_.GetAttribute('Language') -match "(^|;)$language(;|$)"
} | Select-Object -First 1
if (-not $voice) {
    $voice = $englishVoices | Where-Object {
        $_.GetAttribute('Gender') -eq $gender
    } | Select-Object -First 1
}
if (-not $voice) { throw "No Windows English $gender speech voice is installed." }
$synth.Voice = $voice
$stream = New-Object -ComObject SAPI.SpFileStream
$stream.Open($env:TOIC_TTS_OUTPUT, 3, $false)
$synth.AudioOutputStream = $stream
$synth.Speak($text) | Out-Null
$stream.Close()
[Console]::Out.Write($voice.GetDescription())
"""
        environment = os.environ.copy()
        environment["TOIC_TTS_TEXT"] = base64.b64encode(
            article_text.encode("utf-8")
        ).decode("ascii")
        environment["TOIC_TTS_VOICE"] = voice_code
        environment["TOIC_TTS_OUTPUT"] = temp_wav_path

        try:
            result = subprocess.run(
                [
                    "powershell.exe",
                    "-NoProfile",
                    "-NonInteractive",
                    "-Command",
                    script
                ],
                env=environment,
                capture_output=True,
                text=True,
                timeout=120,
                check=True,
                creationflags=subprocess.CREATE_NO_WINDOW
            )
            if os.path.getsize(temp_wav_path) <= 44:
                raise RuntimeError("Windows 語音服務沒有產生有效的 WAV 音訊。")
            os.replace(temp_wav_path, wav_path)
            voice_name = result.stdout.strip()
            self.logger.log(f"   已使用 Windows 離線語音: {voice_name}")
            return wav_path
        except (OSError, subprocess.SubprocessError, RuntimeError) as fallback_error:
            raise RuntimeError(
                f"Edge TTS 無法產生音訊 ({tts_error})，Windows 離線語音備援也失敗: "
                f"{fallback_error}"
            ) from fallback_error
        finally:
            if os.path.exists(temp_wav_path):
                os.remove(temp_wav_path)


if __name__ == "__main__":
    root = tk.Tk()
    app = FlashcardApp(root)
    root.mainloop()